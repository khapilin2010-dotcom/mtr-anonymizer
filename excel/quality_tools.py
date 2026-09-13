"""Quality-control helpers for the field Excel module.

The functions in this module are deliberately UI-agnostic so they can be used
from the desktop quality center, command-line checks and tests.
"""
from collections import Counter, defaultdict
from difflib import SequenceMatcher
import json
from pathlib import Path

from excel.excel_engine import Anonymizer
from excel.expert_engine import protection_losses
from excel.knowledge import case_key, normalize


def previous_decision(row, snapshot):
    """Return exact accumulated knowledge for a row, if it exists."""
    key = case_key(row.get('code', ''), row.get('source', ''), row.get('factory', ''))
    entry = snapshot.entries.get(key)
    if not entry:
        return None
    return {
        'key': key,
        'status': entry.get('status', ''),
        'value': entry.get('value', ''),
        'confirmations': entry.get('confirmations', 0),
        'opposition': entry.get('opposition', 0),
        'score': entry.get('score'),
    }


def removed_ratio(source, final):
    """Approximate share of source characters that disappeared from the result."""
    source = str(source or '')
    final = str(final or '')
    if not source:
        return 0.0
    matcher = SequenceMatcher(None, source, final, autojunk=False)
    kept = sum(block.size for block in matcher.get_matching_blocks())
    return max(0.0, min(1.0, 1.0 - kept / max(1, len(source))))


def row_risks(row, snapshot=None, base=None):
    """Return actionable reasons why a generated row deserves attention."""
    base = base or Anonymizer()
    source = str(row.get('source', '') or '')
    final = str(row.get('automatic', '') or '')
    code = str(row.get('code', '') or '')
    factory = str(row.get('factory', '') or '')
    risks = []

    if source.strip() and not final.strip():
        risks.append('Результат пустой')
    losses = protection_losses(source, final, base, code, factory)
    if losses:
        risks.append('Потеря защищённых признаков: ' + ', '.join(losses))
    ratio = removed_ratio(source, final)
    if ratio >= 0.5:
        risks.append(f'Удалено около {ratio * 100:.0f}% исходного текста')
    if row.get('status') == 'КРАСНЫЙ':
        risks.append('Красный статус программы')

    if snapshot is not None:
        previous = previous_decision(row, snapshot)
        if previous:
            if previous['status'] == 'DISPUTED':
                risks.append('В базе есть противоречащие решения инженеров')
            elif previous['status'] in ('ACTIVE', 'TRUSTED') and normalize(previous['value']) != normalize(final):
                risks.append('Текущий результат отличается от ранее подтверждённого решения')
    return list(dict.fromkeys(risks))


def is_unknown_row(row, snapshot):
    """Rows without an exact accumulated case are best candidates for training."""
    previous = previous_decision(row, snapshot)
    if not previous:
        return True
    return previous['status'] in ('CANDIDATE', 'DISPUTED', 'DISABLED')


def same_case(a, b):
    """Safe in-session grouping for a mass engineer decision.

    We require identical normalized source, factory and current automatic result.
    Code may differ because historical exports sometimes duplicate one resource
    under several legacy code spellings; the user still gets a confirmation count.
    """
    return (
        normalize(a.get('source', '')) == normalize(b.get('source', ''))
        and normalize(a.get('factory', '')) == normalize(b.get('factory', ''))
        and normalize(a.get('automatic', '')) == normalize(b.get('automatic', ''))
    )


def knowledge_audit(snapshot):
    """Find contradictions, weak/unused knowledge and suspicious patterns."""
    issues = []
    fragment_groups = defaultdict(list)
    for key, entry in snapshot.entries.items():
        sample = entry.get('sample', {})
        example = sample.get('source', sample.get('example', ''))
        base = dict(key=key, kind=entry.get('kind', ''), status=entry.get('status', ''),
                    value=entry.get('value', ''), confirmations=entry.get('confirmations', 0),
                    opposition=entry.get('opposition', 0), applications=entry.get('applications', 0),
                    example=example)
        if entry.get('status') == 'DISPUTED':
            issues.append(dict(base, severity='КРАСНЫЙ', issue='Противоречащие решения'))
        elif entry.get('opposition', 0):
            issues.append(dict(base, severity='ЖЁЛТЫЙ', issue='Есть противоположные голоса'))
        if entry.get('kind') == 'rule':
            fragment = str(sample.get('fragment', '')).strip()
            if fragment:
                fragment_groups[' '.join(fragment.casefold().split())].append(entry)
            if entry.get('applications', 0) == 0 and entry.get('status') not in ('DISABLED',):
                issues.append(dict(base, severity='ЖЁЛТЫЙ', issue='Правило пока ни разу не применялось'))
        if entry.get('status') == 'CANDIDATE' and entry.get('confirmations', 0) == 1:
            issues.append(dict(base, severity='ИНФО', issue='Пока только одно независимое подтверждение'))

    for fragment, entries in fragment_groups.items():
        values = {e.get('value') for e in entries if e.get('status') != 'DISABLED'}
        if len(values) > 1:
            first = entries[0]
            sample = first.get('sample', {})
            issues.append({
                'key': first.get('key', ''), 'kind': 'rule', 'status': 'DISPUTED',
                'value': ' / '.join(sorted(v for v in values if v)),
                'confirmations': sum(e.get('confirmations', 0) for e in entries),
                'opposition': sum(e.get('opposition', 0) for e in entries),
                'applications': sum(e.get('applications', 0) for e in entries),
                'example': sample.get('example', ''), 'severity': 'КРАСНЫЙ',
                'issue': 'Один и тот же фрагмент имеет разные правила: ' + fragment,
            })
    order = {'КРАСНЫЙ': 0, 'ЖЁЛТЫЙ': 1, 'ИНФО': 2}
    return sorted(issues, key=lambda x: (order.get(x['severity'], 9), x['issue'], x['key']))


def session_summary(store):
    """Return per-session row statistics for management reporting."""
    result = []
    with store.db() as db:
        sessions = db.execute('SELECT session, count(*) FROM rows GROUP BY session ORDER BY max(rowid) DESC').fetchall()
        for sid, count in sessions:
            bodies = [json.loads(r[0]) for r in db.execute('SELECT body FROM rows WHERE session=?', (sid,))]
            statuses = Counter(r.get('status', '') for r in bodies)
            changed = sum(normalize(r.get('source', '')) != normalize(r.get('automatic', '')) for r in bodies)
            dangerous = sum(bool(row_risks(r, None)) for r in bodies)
            meta = store.metadata(sid)
            result.append({
                'session': sid,
                'file': meta.get('file', ''),
                'timestamp': meta.get('timestamp', ''),
                'rows': count,
                'changed': changed,
                'green': statuses.get('ЗЕЛЁНЫЙ', 0),
                'yellow': statuses.get('ЖЁЛТЫЙ', 0),
                'red': statuses.get('КРАСНЫЙ', 0),
                'dangerous': dangerous,
            })
    return result


def export_manager_report(path, store):
    """Create a compact XLSX report for a supervisor."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    path = Path(path)
    snapshot = store.snapshot()
    sessions = session_summary(store)
    audit = knowledge_audit(snapshot)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Сводка'
    headers = ['Показатель', 'Значение']
    ws.append(headers)
    total_rows = sum(s['rows'] for s in sessions)
    total_changed = sum(s['changed'] for s in sessions)
    rows = [
        ('Обработано сеансов', len(sessions)),
        ('Обработано строк', total_rows),
        ('Изменено строк', total_changed),
        ('Зелёных строк', sum(s['green'] for s in sessions)),
        ('Жёлтых строк', sum(s['yellow'] for s in sessions)),
        ('Красных строк', sum(s['red'] for s in sessions)),
        ('Потенциально опасных строк', sum(s['dangerous'] for s in sessions)),
        ('Записей базы знаний', len(snapshot.entries)),
        ('CANDIDATE', sum(e.get('status') == 'CANDIDATE' for e in snapshot.entries.values())),
        ('ACTIVE', sum(e.get('status') == 'ACTIVE' for e in snapshot.entries.values())),
        ('TRUSTED', sum(e.get('status') == 'TRUSTED' for e in snapshot.entries.values())),
        ('DISPUTED', sum(e.get('status') == 'DISPUTED' for e in snapshot.entries.values())),
        ('Замечаний контроля базы', len(audit)),
    ]
    for row in rows:
        ws.append(row)

    sessions_ws = wb.create_sheet('Сеансы')
    sessions_ws.append(['Дата', 'Файл', 'Строк', 'Изменено', 'Зелёных', 'Жёлтых', 'Красных', 'Опасных'])
    for s in sessions:
        sessions_ws.append([s['timestamp'], s['file'], s['rows'], s['changed'], s['green'], s['yellow'], s['red'], s['dangerous']])

    audit_ws = wb.create_sheet('Контроль базы')
    audit_ws.append(['Серьёзность', 'Проблема', 'Вид', 'Статус', 'Решение', 'Подтверждений', 'Против', 'Применений', 'Пример', 'Ключ'])
    for item in audit:
        audit_ws.append([item['severity'], item['issue'], item['kind'], item['status'], item['value'],
                         item['confirmations'], item['opposition'], item['applications'], item['example'], item['key']])

    factories = Counter()
    with store.db() as db:
        for body, in db.execute('SELECT body FROM rows'):
            row = json.loads(body)
            factory = str(row.get('detected_factory') or row.get('factory') or '').strip()
            if factory:
                factories[factory] += 1
    fac_ws = wb.create_sheet('Производители')
    fac_ws.append(['Производитель', 'Строк'])
    for factory, count in factories.most_common(50):
        fac_ws.append([factory, count])

    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1767A6')
            cell.alignment = Alignment(wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(60, max(12, max(len(str(c.value or '')) for c in column[:200]) + 2))
            sheet.column_dimensions[letter].width = width
    wb.save(path)
    wb.close()
    return path
