"""Import previously reviewed/anonymized spreadsheets into expert knowledge.

This is intentionally stricter than ordinary processing: the workbook must
contain both the source and the reviewed anonymized name.  Nothing is inferred
from untouched arbitrary spreadsheets.
"""
from pathlib import Path
import csv
import io
import uuid

from excel.excel_engine import Anonymizer
from excel.expert_engine import protection_losses
from excel.knowledge import case_key, normalize


ALIASES = {
    'code': {'код autodocs', 'код автодокс', 'код', 'autodocs', 'код мтр', 'код ресурса'},
    'source': {'исходное наименование', 'наименование исходное', 'наименование',
               'наименование мтр', 'наименование и техническая характеристика'},
    'factory': {'завод', 'изготовитель', 'производитель', 'поставщик',
                'завод/изготовитель/поставщик'},
    'final': {'обезличенное наименование', 'обезличенное', 'итоговое наименование',
              'наименование обезличенное'},
}


def _norm(value):
    return ' '.join(str(value or '').strip().lower().replace('ё', 'е').split())


def _header(values):
    mapping = {}
    for index, value in enumerate(values):
        key = _norm(value)
        for name, aliases in ALIASES.items():
            if key in aliases:
                mapping.setdefault(name, index)
    if 'source' in mapping and 'final' in mapping:
        return mapping
    return None


def _iter_rows(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False, keep_vba=suffix == '.xlsm')
        try:
            for ws in wb.worksheets:
                mapping = None
                for row_no, cells in enumerate(ws.iter_rows(), 1):
                    values = [c.value for c in cells]
                    if mapping is None:
                        if row_no > 30:
                            break
                        mapping = _header(values)
                        if mapping:
                            continue
                    else:
                        yield ws.title, row_no, mapping, values, {
                            i for i, c in enumerate(cells) if c.data_type in ('f', 'e')
                        }
        finally:
            wb.close()
        return
    if suffix == '.xls':
        import xlrd
        wb = xlrd.open_workbook(str(path))
        try:
            for ws in wb.sheets():
                mapping = None
                for row_no in range(ws.nrows):
                    values = ws.row_values(row_no)
                    if mapping is None:
                        if row_no >= 30:
                            break
                        mapping = _header(values)
                        if mapping:
                            continue
                    else:
                        yield ws.name, row_no + 1, mapping, values, set()
        finally:
            wb.release_resources()
        return
    if suffix == '.csv':
        raw = path.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1251')
        for delimiter in (';', '\t', ','):
            rows = list(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter))
            mapping = None
            header_index = None
            for i, row in enumerate(rows[:30]):
                mapping = _header(row)
                if mapping:
                    header_index = i
                    break
            if mapping is not None:
                for i, row in enumerate(rows[header_index + 1:], header_index + 2):
                    yield 'CSV', i, mapping, row, set()
                return
        raise ValueError('Не найдены заголовки старой проверенной выборки.')
    raise ValueError('Поддерживаются XLSX, XLSM, XLS и CSV.')


def preview_legacy(path, store):
    """Return a plan without writing any decision."""
    store.sync()
    snapshot = store.snapshot()
    base = Anonymizer()
    plan = []
    issues = []
    repeated = 0
    conflicts = 0
    seen = 0
    for sheet, row_no, cols, values, formulas in _iter_rows(path):
        seen += 1
        def get(name):
            i = cols.get(name)
            return '' if i is None or i >= len(values) or values[i] is None else str(values[i])
        important = {cols['source'], cols['final']}
        if formulas & important:
            issues.append((sheet, row_no, 'Формула в исходном или обезличенном наименовании'))
            continue
        source, final = get('source').strip(), get('final').strip()
        if not source and not final:
            continue
        if not source or not final:
            issues.append((sheet, row_no, 'Нужны и исходное, и обезличенное наименование'))
            continue
        code, factory = get('code').strip(), get('factory').strip()
        losses = protection_losses(source, final, base, code, factory)
        if losses:
            issues.append((sheet, row_no, 'Итог теряет защищённые признаки: ' + ', '.join(losses)))
            continue
        key = case_key(code, source, factory)
        current = snapshot.entries.get(key)
        if current and normalize(current.get('value', '')) == normalize(final):
            repeated += 1
            continue
        conflict = bool(current and current.get('status') != 'DISABLED'
                        and normalize(current.get('value', '')) != normalize(final))
        conflicts += int(conflict)
        plan.append(dict(sheet=sheet, row=row_no, code=code, source=source,
                         factory=factory, final=final, key=key, conflict=conflict,
                         current=current.get('value', '') if current else ''))
    if seen == 0:
        raise ValueError('В файле не найдена таблица с исходным и обезличенным наименованием.')
    return dict(rows=seen, ready=len(plan), repeated=repeated, conflicts=conflicts,
                issues=issues, plan=plan)


def commit_legacy(path, store, accept_conflicts=False):
    report = preview_legacy(path, store)
    accepted = 0
    skipped_conflicts = 0
    batch = 'legacy-import:' + uuid.uuid4().hex
    for item in report['plan']:
        if item['conflict'] and not accept_conflicts:
            skipped_conflicts += 1
            continue
        row = dict(id=str(uuid.uuid4()), session=batch, batch_id=batch,
                   code=item['code'], source=item['source'], factory=item['factory'],
                   automatic=item['source'], classification=None,
                   explicit_feedback=True,
                   provenance={'source': 'Импорт старой проверенной выборки',
                               'file': str(Path(path).name),
                               'sheet': item['sheet'], 'row': item['row']})
        event = store.decide(row, item['final'], 'Импорт проверенного Excel')
        accepted += int(event is not None)
    store.sync()
    return dict(accepted=accepted, repeated=report['repeated'],
                conflicts=report['conflicts'], skipped_conflicts=skipped_conflicts,
                issues=report['issues'])
