"""Spreadsheet I/O. Original cells remain intact; audit fields are appended."""
import csv
import io
import os
from pathlib import Path
import tempfile

from excel.excel_engine import Anonymizer

HEADERS = ('Выявленный завод/производитель', 'Обезличенное наименование',
           'Статус проверки', 'Что именно удалено')
SUPPORTED = {'.xlsx', '.xlsm', '.xls', '.csv'}
HEADER_NAMES = {
    'code': {'код autodocs', 'код автодокс', 'код', 'autodocs', 'код мтр', 'код ресурса',
             '№ смет/ код ресурса', '№ смет / код ресурса'},
    'name': {'наименование', 'наименование мтр', 'мтр', 'описание',
             'наименование и техническая характеристика', 'исходное наименование'},
    'factory': {'завод', 'изготовитель', 'производитель', 'поставщик', 'завод/изготовитель/поставщик'},
}


def norm_header(value):
    return ' '.join(str(value or '').lower().replace('ё', 'е').split())


def find_header(rows):
    for index, row in enumerate(rows):
        mapping = {}
        for col, value in enumerate(row):
            for key, names in HEADER_NAMES.items():
                if norm_header(value) in names:
                    mapping.setdefault(key, col)
        if 'name' in mapping and ('code' in mapping or 'factory' in mapping):
            return index, mapping
    raise ValueError('Не найдены столбцы «Наименование» и «Код Автодокс/Код ресурса» или «Производитель».')


def select_titles(titles):
    keys = {t: norm_header(t.replace('_', ' ')) for t in titles}
    for preferred in ('готово', 'выборка оборудования'):
        chosen = [t for t, key in keys.items() if key == preferred]
        if chosen:
            return chosen
    return list(titles)


def new_report():
    return dict(rows=0, changed=0, green=0, yellow=0, sheets=0, skipped_formulas=0,
                skipped_sheets=[], warnings=[])


def record(result, report):
    report['rows'] += 1
    report['changed'] += int(result['changed'])
    report['green' if result['status'] == 'ЗЕЛЁНЫЙ' else 'yellow'] += 1
    # Include the reason in the status cell so a yellow row is actionable.
    status = result['status'] + (': ' + result['reason'] if result['reason'] else '')
    return [result['factory'], result['text'], status, '; '.join(result['removed'])]


def text_value(value):
    return '' if value is None else str(value)


def process_excel(src, dst, az, progress=None):
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    wb = load_workbook(src, keep_vba=src.suffix.lower() == '.xlsm')
    report = new_report()
    try:
        for title in select_titles(wb.sheetnames):
            ws = wb[title]
            try:
                header, cols = find_header(ws.iter_rows(min_row=1, max_row=min(30, ws.max_row), values_only=True))
            except ValueError:
                report['skipped_sheets'].append(title)
                continue
            report['sheets'] += 1
            offset = ws.max_column
            for col, label in enumerate(HEADERS, offset + 1):
                cell = ws.cell(header + 1, col, label)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1767A6')
                cell.alignment = Alignment(wrap_text=True)
                ws.column_dimensions[cell.column_letter].width = 55 if col != offset + 3 else 36
            for row in range(header + 2, ws.max_row + 1):
                source = ws.cell(row, cols['name'] + 1)
                name = text_value(source.value)
                if not name.strip() or source.data_type == 'e':
                    continue
                code_cell = ws.cell(row, cols['code'] + 1) if 'code' in cols else None
                factory_cell = ws.cell(row, cols['factory'] + 1) if 'factory' in cols else None
                if any(c is not None and c.data_type == 'f' for c in (source, code_cell, factory_cell)):
                    report['skipped_formulas'] += 1
                    ws.cell(row, offset + 3, 'ЖЁЛТЫЙ: формула в исходных данных; требуется проверка значений')
                    continue
                code = text_value(code_cell.value) if code_cell is not None else ''
                factory = text_value(factory_cell.value) if factory_cell is not None else ''
                if row <= header + 6 and code in {'1', '2'} and name in {'3', '4'}:
                    continue
                result = az.anonymize(name, code, factory)
                for col, value in enumerate(record(result, report), offset + 1):
                    cell = ws.cell(row, col, value)
                    cell.data_type = 's'  # Never turn a generated value into a formula.
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                ws.cell(row, offset + 3).fill = PatternFill(
                    'solid', fgColor='D9EAD3' if result['status'] == 'ЗЕЛЁНЫЙ' else 'FFF2CC')
                if progress and (row % 100 == 0 or row == ws.max_row):
                    progress(f'{title}: {row - header - 1} / {ws.max_row - header - 1}')
        if not report['sheets']:
            raise ValueError('Ни на одном рабочем листе не найдены столбцы МТР.')
        wb.save(dst)
    finally:
        wb.close()
    return report


def process_csv(src, dst, az, progress=None):
    raw = src.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('cp1251')
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=';,\t')
    except csv.Error:
        # Preamble rows can confuse Sniffer; locate a recognizable header.
        dialect = None
        for delimiter in (';', '\t', ','):
            candidate = list(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter))
            try:
                find_header(candidate[:30])
                dialect = delimiter
                break
            except ValueError:
                pass
        if dialect is None:
            raise ValueError('Не удалось определить разделитель и заголовки CSV.')
    options = {'delimiter': dialect} if isinstance(dialect, str) else {'dialect': dialect}
    rows = list(csv.reader(io.StringIO(text, newline=''), **options))
    header, cols = find_header(rows[:30])
    offset = max(map(len, rows))
    report = new_report()
    report['sheets'] = 1
    rows[header].extend([''] * (offset - len(rows[header])))
    rows[header].extend(HEADERS)
    for index in range(header + 1, len(rows)):
        row = rows[index]
        row.extend([''] * (offset - len(row)))
        name = row[cols['name']]
        if not name.strip():
            continue
        result = az.anonymize(name, row[cols['code']] if 'code' in cols else '',
                              row[cols['factory']] if 'factory' in cols else '')
        row.extend(record(result, report))
        if progress and index % 100 == 0:
            progress(f'CSV: {index - header} / {len(rows) - header - 1}')
    with open(dst, 'w', encoding='utf-8-sig', newline='') as handle:
        csv.writer(handle, **options).writerows(rows)
    return report


def process_xls(src, dst, az, progress=None):
    import xlrd
    import xlwt
    from xlutils.copy import copy as copy_xls
    source = xlrd.open_workbook(str(src), formatting_info=True)
    # xlutils cannot retain BIFF formula expressions. Refuse silent conversion.
    # Scan actual BIFF record headers, not arbitrary byte substrings.
    from xlrd.compdoc import CompDoc
    import struct
    stream = CompDoc(src.read_bytes()).get_named_stream('Workbook')
    if stream is None:
        stream = CompDoc(src.read_bytes()).get_named_stream('Book')
    position = 0
    while stream and position + 4 <= len(stream):
        kind, length = struct.unpack_from('<HH', stream, position)
        if kind in (0x0006, 0x0206, 0x0406):
            raise ValueError('XLS содержит формулы. Сохраните исходник в Excel как XLSX и повторите обработку: формулы будут сохранены.')
        position += 4 + length
    output = copy_xls(source)
    report = new_report()
    for title in select_titles(source.sheet_names()):
        sheet = source.sheet_by_name(title)
        try:
            header, cols = find_header([sheet.row_values(i) for i in range(min(30, sheet.nrows))])
        except ValueError:
            report['skipped_sheets'].append(title)
            continue
        if sheet.ncols + 4 > 256:
            raise ValueError('В XLS недостаточно места для четырёх столбцов. Сохраните исходник как XLSX.')
        report['sheets'] += 1
        writable = output.get_sheet(source.sheet_names().index(title))
        for col, label in enumerate(HEADERS, sheet.ncols):
            writable.write(header, col, label)
            writable.col(col).width = 14000
        for row in range(header + 1, sheet.nrows):
            if sheet.cell_type(row, cols['name']) == xlrd.XL_CELL_ERROR:
                continue
            name = text_value(sheet.cell_value(row, cols['name']))
            if not name.strip():
                continue
            result = az.anonymize(name, text_value(sheet.cell_value(row, cols['code'])) if 'code' in cols else '',
                                  text_value(sheet.cell_value(row, cols['factory'])) if 'factory' in cols else '')
            for col, value in enumerate(record(result, report), sheet.ncols):
                writable.write(row, col, value)
            if progress and row % 100 == 0:
                progress(f'{title}: {row - header} / {sheet.nrows - header - 1}')
    if not report['sheets']:
        raise ValueError('Не найдены столбцы МТР в XLS.')
    output.save(str(dst))
    source.release_resources()
    return report


def process_file(src, output_dir, az=None, progress=None):
    src, output_dir = Path(src), Path(output_dir)
    if src.suffix.lower() not in SUPPORTED:
        raise ValueError('Поддерживаются XLSX, XLSM, XLS и CSV.')
    output_dir.mkdir(parents=True, exist_ok=True)
    suffix = src.suffix.lower()
    # Unique names prevent overwriting any source or previous result.
    index = 0
    while True:
        tail = f'_{index}' if index else ''
        dst = output_dir / f'{src.stem}_обезличено{tail}{suffix}'
        try:
            reservation = os.open(dst, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.close(reservation)
            break
        except FileExistsError:
            index += 1
    fd, temporary = tempfile.mkstemp(prefix='.mtr-', suffix=suffix, dir=output_dir)
    os.close(fd)
    try:
        processor = {'.csv': process_csv, '.xls': process_xls}.get(suffix, process_excel)
        report = processor(src, Path(temporary), az or Anonymizer(), progress)
        os.replace(temporary, dst)
        return dst, report
    except BaseException:
        Path(temporary).unlink(missing_ok=True)
        dst.unlink(missing_ok=True)
        raise
