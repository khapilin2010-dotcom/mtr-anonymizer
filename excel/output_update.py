"""Apply engineer decisions back to the already generated workbook.

The source workbook is never touched.  This module only updates the generated
output registered for a processing session, so the simple review UI can finish
with one ready-to-use Excel file instead of asking the engineer to process the
source again.
"""
import csv
import io
import os
from pathlib import Path
import tempfile

from excel.file_io import HEADERS
from excel.review_io import REVIEW_HEADERS
from excel.knowledge import fragments


def _removed(source, final):
    deleted, _ = fragments(str(source or ''), str(final or ''))
    return '; '.join(x.strip() for x in deleted if x.strip())


def _header_map(values):
    return {str(value or '').strip(): index for index, value in enumerate(values)}


def _atomic_target(path):
    path = Path(path)
    fd, temporary = tempfile.mkstemp(prefix='.mtr-review-', suffix=path.suffix, dir=path.parent)
    os.close(fd)
    return Path(temporary)


def _xlsx(path, decisions):
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    keep_vba = path.suffix.lower() == '.xlsm'
    wb = load_workbook(path, keep_vba=keep_vba)
    temporary = _atomic_target(path)
    try:
        by_sheet = {}
        for decision in decisions:
            by_sheet.setdefault(decision['row'].get('sheet', ''), []).append(decision)
        for sheet_name, items in by_sheet.items():
            if sheet_name not in wb.sheetnames:
                raise ValueError(f'В результате не найден лист «{sheet_name}».')
            ws = wb[sheet_name]
            header_row = None
            mapping = None
            for number, cells in enumerate(ws.iter_rows(min_row=1, max_row=min(30, ws.max_row), values_only=True), 1):
                values = [str(x or '').strip() for x in cells]
                if HEADERS[1] in values:
                    header_row = number
                    mapping = _header_map(values)
                    break
            if not mapping:
                raise ValueError(f'На листе «{sheet_name}» не найдены столбцы результата.')
            for required in (HEADERS[1], HEADERS[2], HEADERS[3]):
                if required not in mapping:
                    raise ValueError(f'На листе «{sheet_name}» нет столбца «{required}».')
            for item in items:
                row = item['row']
                excel_row = int(row['row'])
                final = str(item['final'])
                action = str(item['action'])
                ws.cell(excel_row, mapping[HEADERS[1]] + 1, final).data_type = 's'
                status_cell = ws.cell(excel_row, mapping[HEADERS[2]] + 1,
                                      'ПРОВЕРЕНО ИНЖЕНЕРОМ: ' + action)
                status_cell.data_type = 's'
                status_cell.fill = PatternFill('solid', fgColor='D9EAD3')
                ws.cell(excel_row, mapping[HEADERS[3]] + 1,
                        _removed(row.get('source', ''), final)).data_type = 's'
                if REVIEW_HEADERS[0] in mapping:
                    ws.cell(excel_row, mapping[REVIEW_HEADERS[0]] + 1, action).data_type = 's'
        wb.save(temporary)
        wb.close()
        os.replace(temporary, path)
    finally:
        try:
            wb.close()
        except Exception:
            pass
        temporary.unlink(missing_ok=True)


def _csv(path, decisions):
    raw = path.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
        encoding = 'utf-8-sig'
    except UnicodeDecodeError:
        text = raw.decode('cp1251')
        encoding = 'cp1251'
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=';,\t')
        options = {'dialect': dialect}
    except csv.Error:
        delimiter = next((d for d in (';', '\t', ',') if HEADERS[1] in next(csv.reader(io.StringIO(text), delimiter=d), [])), ';')
        options = {'delimiter': delimiter}
    rows = list(csv.reader(io.StringIO(text, newline=''), **options))
    header_index = next((i for i, row in enumerate(rows[:30]) if HEADERS[1] in row), None)
    if header_index is None:
        raise ValueError('В CSV не найдены столбцы результата.')
    mapping = _header_map(rows[header_index])
    for item in decisions:
        row = item['row']
        index = int(row['row']) - 1
        if index < 0 or index >= len(rows):
            raise ValueError('Не найдена строка результата CSV: ' + str(row.get('row')))
        values = rows[index]
        if len(values) < len(rows[header_index]):
            values.extend([''] * (len(rows[header_index]) - len(values)))
        final = str(item['final'])
        action = str(item['action'])
        values[mapping[HEADERS[1]]] = final
        values[mapping[HEADERS[2]]] = 'ПРОВЕРЕНО ИНЖЕНЕРОМ: ' + action
        values[mapping[HEADERS[3]]] = _removed(row.get('source', ''), final)
        if REVIEW_HEADERS[0] in mapping:
            values[mapping[REVIEW_HEADERS[0]]] = action
    temporary = _atomic_target(path)
    try:
        with temporary.open('w', encoding=encoding, newline='') as handle:
            csv.writer(handle, **options).writerows(rows)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _xls(path, decisions):
    import xlrd
    from xlutils.copy import copy as copy_xls

    book = xlrd.open_workbook(str(path), formatting_info=True)
    writable = copy_xls(book)
    by_sheet = {}
    for decision in decisions:
        by_sheet.setdefault(decision['row'].get('sheet', ''), []).append(decision)
    for sheet_name, items in by_sheet.items():
        if sheet_name not in book.sheet_names():
            raise ValueError(f'В результате не найден лист «{sheet_name}».')
        index = book.sheet_names().index(sheet_name)
        source_sheet = book.sheet_by_index(index)
        target_sheet = writable.get_sheet(index)
        header_index = next((i for i in range(min(30, source_sheet.nrows))
                             if HEADERS[1] in list(map(str, source_sheet.row_values(i)))), None)
        if header_index is None:
            raise ValueError(f'На листе «{sheet_name}» не найдены столбцы результата.')
        mapping = _header_map(list(map(str, source_sheet.row_values(header_index))))
        for item in items:
            row = item['row']
            excel_row = int(row['row']) - 1
            final = str(item['final'])
            action = str(item['action'])
            target_sheet.write(excel_row, mapping[HEADERS[1]], final)
            target_sheet.write(excel_row, mapping[HEADERS[2]], 'ПРОВЕРЕНО ИНЖЕНЕРОМ: ' + action)
            target_sheet.write(excel_row, mapping[HEADERS[3]], _removed(row.get('source', ''), final))
            if REVIEW_HEADERS[0] in mapping:
                target_sheet.write(excel_row, mapping[REVIEW_HEADERS[0]], action)
    temporary = _atomic_target(path)
    try:
        writable.save(str(temporary))
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
        book.release_resources()


def apply_decisions(path, decisions):
    """Apply a batch of decisions to one generated output file atomically."""
    path = Path(path)
    if not decisions:
        return path
    suffix = path.suffix.lower()
    if suffix in ('.xlsx', '.xlsm'):
        _xlsx(path, decisions)
    elif suffix == '.csv':
        _csv(path, decisions)
    elif suffix == '.xls':
        _xls(path, decisions)
    else:
        raise ValueError('Неподдерживаемый формат результата: ' + suffix)
    return path
