"""Import old reviewed workbooks with user-selected column mapping."""
from pathlib import Path
import csv
import io
import uuid

from excel.excel_engine import Anonymizer
from excel.expert_engine import protection_losses
from excel.knowledge import case_key, normalize

NONE = '— не использовать —'


def _text(value):
    return '' if value is None else str(value)


def _csv_rows(path):
    raw = Path(path).read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('cp1251')
    best = None
    for delimiter in (';', '\t', ','):
        rows = list(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter))
        score = max((sum(bool(str(v).strip()) for v in row) for row in rows[:30]), default=0)
        if best is None or score > best[0]:
            best = (score, delimiter, rows)
    return best[2] if best else []


def header_candidates(path):
    """Return candidate header rows from the first 30 rows of each sheet."""
    path = Path(path)
    suffix = path.suffix.lower()
    result = []
    if suffix in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            for ws in wb.worksheets:
                for row_no, cells in enumerate(ws.iter_rows(max_row=min(30, ws.max_row)), 1):
                    headers = [_text(c.value).strip() for c in cells]
                    if sum(bool(x) for x in headers) >= 2:
                        result.append({'sheet': ws.title, 'row': row_no, 'headers': headers})
        finally:
            wb.close()
    elif suffix == '.xls':
        import xlrd
        wb = xlrd.open_workbook(str(path))
        try:
            for ws in wb.sheets():
                for i in range(min(30, ws.nrows)):
                    headers = [_text(x).strip() for x in ws.row_values(i)]
                    if sum(bool(x) for x in headers) >= 2:
                        result.append({'sheet': ws.name, 'row': i + 1, 'headers': headers})
        finally:
            wb.release_resources()
    elif suffix == '.csv':
        rows = _csv_rows(path)
        for i, row in enumerate(rows[:30], 1):
            headers = [_text(x).strip() for x in row]
            if sum(bool(x) for x in headers) >= 2:
                result.append({'sheet': 'CSV', 'row': i, 'headers': headers})
    else:
        raise ValueError('Поддерживаются XLSX, XLSM, XLS и CSV.')
    if not result:
        raise ValueError('В первых 30 строках не найдено подходящей строки заголовков.')
    return result


def guess_mapping(headers):
    """Best-effort guesses; the engineer can always override them."""
    aliases = {
        'code': ('код автодокс', 'код autodocs', 'код мтр', 'код ресурса', 'код'),
        'source': ('исходное наименование', 'наименование исходное', 'наименование мтр', 'наименование'),
        'factory': ('завод', 'изготовитель', 'производитель', 'поставщик'),
        'final': ('обезличенное наименование', 'обезличенное', 'итоговое наименование'),
    }
    normalized = [' '.join(str(x).casefold().replace('ё', 'е').split()) for x in headers]
    result = {'code': None, 'source': None, 'factory': None, 'final': None}
    # Match semantic names before generic "наименование", irrespective of column order.
    used = set()
    for key in ('final', 'source', 'code', 'factory'):
        names = aliases[key]
        matches = [(int(value != name), rank, i) for rank, name in enumerate(names)
                   for i, value in enumerate(normalized) if i not in used and name in value]
        if matches:
            result[key] = min(matches)[2]
            used.add(result[key])
    return result


def _iter_mapped(path, candidate, mapping):
    path = Path(path)
    suffix = path.suffix.lower()
    sheet_name, header_row = candidate['sheet'], int(candidate['row'])
    if suffix in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError('Выбранный лист больше не найден.')
            ws = wb[sheet_name]
            for row_no, cells in enumerate(ws.iter_rows(min_row=header_row + 1), header_row + 1):
                values = [c.value for c in cells]
                formulas = {i for i, c in enumerate(cells) if c.data_type in ('f', 'e')}
                yield sheet_name, row_no, values, formulas
        finally:
            wb.close()
        return
    if suffix == '.xls':
        import xlrd
        wb = xlrd.open_workbook(str(path))
        try:
            ws = wb.sheet_by_name(sheet_name)
            for i in range(header_row, ws.nrows):
                yield sheet_name, i + 1, ws.row_values(i), set()
        finally:
            wb.release_resources()
        return
    rows = _csv_rows(path)
    for i, row in enumerate(rows[header_row:], header_row + 1):
        yield 'CSV', i, row, set()


def preview_mapped(path, store, candidate, mapping):
    """Validate a manual column mapping without writing to the knowledge base."""
    if mapping.get('source') is None or mapping.get('final') is None:
        raise ValueError('Обязательно укажите колонки исходного и обезличенного наименования.')
    if mapping['source'] == mapping['final']:
        raise ValueError('Исходное и обезличенное наименование не могут быть одной колонкой.')
    store.sync()
    snapshot = store.snapshot()
    base = Anonymizer()
    plan, issues = [], []
    repeated = conflicts = seen = 0

    for sheet, row_no, values, formulas in _iter_mapped(path, candidate, mapping):
        seen += 1
        def get(name):
            index = mapping.get(name)
            return '' if index is None or index >= len(values) else _text(values[index]).strip()
        important = {mapping['source'], mapping['final']}
        if formulas & important:
            issues.append((sheet, row_no, 'Формула в исходном или обезличенном наименовании'))
            continue
        source, final = get('source'), get('final')
        if not source and not final:
            continue
        if not source or not final:
            issues.append((sheet, row_no, 'Нужны и исходное, и обезличенное наименование'))
            continue
        code, factory = get('code'), get('factory')
        losses = protection_losses(source, final, base, code, factory)
        if losses:
            issues.append((sheet, row_no, 'Итог теряет защищённые признаки: ' + ', '.join(losses)))
            continue
        key = case_key(code, source, factory)
        current = snapshot.entries.get(key)
        if current and normalize(current.get('value', '')) == normalize(final):
            repeated += 1
            continue
        conflict = bool(current and current.get('status') != 'DISABLED'
                        and normalize(current.get('value', '')) != normalize(final))
        conflicts += int(conflict)
        plan.append({'sheet': sheet, 'row': row_no, 'code': code, 'source': source,
                     'factory': factory, 'final': final, 'key': key, 'conflict': conflict,
                     'current': current.get('value', '') if current else ''})
    if seen == 0:
        raise ValueError('После выбранной строки заголовков нет данных.')
    return {'rows': seen, 'ready': len(plan), 'repeated': repeated, 'conflicts': conflicts,
            'issues': issues, 'plan': plan}


def commit_mapped(path, store, candidate, mapping, accept_conflicts=False):
    report = preview_mapped(path, store, candidate, mapping)
    accepted = skipped_conflicts = 0
    batch = 'mapped-import:' + uuid.uuid4().hex
    for item in report['plan']:
        if item['conflict'] and not accept_conflicts:
            skipped_conflicts += 1
            continue
        row = dict(id=str(uuid.uuid4()), session=batch, batch_id=batch,
                   code=item['code'], source=item['source'], factory=item['factory'],
                   automatic=item['source'], classification=None, explicit_feedback=True,
                   provenance={'source': 'Импорт с ручным сопоставлением колонок',
                               'file': Path(path).name, 'sheet': item['sheet'], 'row': item['row']})
        accepted += int(store.decide(row, item['final'], 'Импорт старого Excel') is not None)
    store.sync()
    return {'accepted': accepted, 'repeated': report['repeated'], 'conflicts': report['conflicts'],
            'skipped_conflicts': skipped_conflicts, 'issues': report['issues']}


def mapping_dialog(parent, store):
    """Interactive mapping wizard used by the launcher."""
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    path = filedialog.askopenfilename(parent=parent, title='Старая проверенная выборка',
                                      filetypes=[('Excel / CSV', '*.xlsx *.xlsm *.xls *.csv')])
    if not path:
        return None
    candidates = header_candidates(path)
    win = tk.Toplevel(parent)
    win.title('Сопоставление колонок старого Excel')
    win.geometry('880x560')
    win.transient(parent)
    win.grab_set()

    ttk.Label(win, text=Path(path).name, font=('Segoe UI', 13, 'bold')).pack(anchor='w', padx=14, pady=(14, 4))
    ttk.Label(win, text='Выберите строку заголовков и укажите, где находятся нужные данные. '
                         'Код Автодокс и завод необязательны.', wraplength=830).pack(anchor='w', padx=14)

    labels = []
    for c in candidates:
        short = ' | '.join(x for x in c['headers'][:5] if x)[:110]
        labels.append(f"{c['sheet']} • строка {c['row']} • {short}")
    selected_header = tk.StringVar(value=labels[0])
    row = ttk.Frame(win); row.pack(fill='x', padx=14, pady=10)
    ttk.Label(row, text='Заголовки:').pack(side='left')
    combo = ttk.Combobox(row, state='readonly', textvariable=selected_header, values=labels, width=100)
    combo.pack(side='left', fill='x', expand=True, padx=(8, 0))

    field_vars = {k: tk.StringVar() for k in ('source', 'final', 'code', 'factory')}
    field_names = [('source', 'Исходное наименование *'), ('final', 'Обезличенное наименование *'),
                   ('code', 'Код Автодокс'), ('factory', 'Завод / производитель')]
    field_boxes = {}
    form = ttk.LabelFrame(win, text='Сопоставление', padding=12); form.pack(fill='x', padx=14, pady=6)
    for r, (key, title) in enumerate(field_names):
        ttk.Label(form, text=title).grid(row=r, column=0, sticky='w', pady=5)
        box = ttk.Combobox(form, state='readonly', textvariable=field_vars[key], width=72)
        box.grid(row=r, column=1, sticky='ew', padx=(12, 0), pady=5)
        field_boxes[key] = box
    form.columnconfigure(1, weight=1)

    def configure_fields(*_):
        index = labels.index(selected_header.get())
        headers = candidates[index]['headers']
        options = [NONE] + [f'{i + 1}: {h or "(пустой заголовок)"}' for i, h in enumerate(headers)]
        guessed = guess_mapping(headers)
        for key, box in field_boxes.items():
            box['values'] = options
            guess = guessed.get(key)
            field_vars[key].set(options[guess + 1] if guess is not None else NONE)
    combo.bind('<<ComboboxSelected>>', configure_fields)
    configure_fields()

    info = tk.StringVar(value='Перед записью программа покажет число новых решений, повторов, конфликтов и проблемных строк.')
    ttk.Label(win, textvariable=info, wraplength=830).pack(anchor='w', padx=14, pady=8)
    result = {'report': None}

    def selected_mapping():
        mapping = {}
        for key, var in field_vars.items():
            value = var.get()
            mapping[key] = None if value == NONE else int(value.split(':', 1)[0]) - 1
        return mapping

    def execute():
        try:
            candidate = candidates[labels.index(selected_header.get())]
            mapping = selected_mapping()
            plan = preview_mapped(path, store, candidate, mapping)
            text = (f"Готово к импорту: {plan['ready']}\nУже есть в базе: {plan['repeated']}\n"
                    f"Конфликтов: {plan['conflicts']}\nПроблемных строк: {len(plan['issues'])}")
            accept_conflicts = False
            if plan['conflicts']:
                accept_conflicts = messagebox.askyesno('Конфликты', text + '\n\nДобавить конфликтующие решения как новые голоса?', parent=win)
            if not messagebox.askyesno('Импорт', text + '\n\nПродолжить?', parent=win):
                return
            report = commit_mapped(path, store, candidate, mapping, accept_conflicts)
            result['report'] = report
            messagebox.showinfo('Готово', f"Принято решений: {report['accepted']}. Проблем: {len(report['issues'])}.", parent=win)
            win.destroy()
        except Exception as exc:
            messagebox.showerror('Не удалось импортировать', str(exc), parent=win)

    buttons = ttk.Frame(win, padding=14); buttons.pack(fill='x', side='bottom')
    ttk.Button(buttons, text='Отмена', command=win.destroy).pack(side='right')
    ttk.Button(buttons, text='Проверить и импортировать', command=execute).pack(side='right', padx=(0, 8))
    parent.wait_window(win)
    return result['report']
