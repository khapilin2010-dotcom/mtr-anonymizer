"""Verified spreadsheet return. Never infer approval from an untouched cell."""
from collections import Counter
import csv
from pathlib import Path

from excel.file_io import HEADERS, find_header, text_value
from excel.knowledge import code_key, normalize

REVIEW_HEADERS = ('Решение инженера', 'Комментарий инженера', 'Источник знания', 'Доверие', 'MTR_ROW_ID', 'MTR_SESSION_ID')
ACTIONS = {'правильно': 'Правильно', 'исправить': 'Исправить', 'оставить как в исходном': 'Оставить как в исходном'}


def review_values(session, sheet, row, code, name, factory, result):
    entry = session.add(sheet, row, code, name, factory, result)
    info = result.get('knowledge', {})
    return ['', '', info.get('source', 'Статическая база'),
            info.get('status', 'LEGACY') + ' / независимых: ' + str(info.get('confirmations', 0)), entry['id'], session.id]


def read_return(path):
    """Yield (sheet, row number, headings, values, formula-column indexes)."""
    path = Path(path)
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            for ws in wb:
                head = None
                for num, cells in enumerate(ws.iter_rows(), 1):
                    values = [text_value(c.value) for c in cells]
                    if head is None:
                        if num > 30:
                            break
                        if HEADERS[1] in values:
                            head = values
                        continue
                    yield ws.title, num, head, values, {i for i, c in enumerate(cells) if c.data_type in ('f', 'e')}
        finally:
            wb.close()
    elif path.suffix.lower() == '.csv':
        raw = path.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1251')
        import io
        for delimiter in (';', '\t', ','):
            rows = list(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter))
            indexes = [i for i, row in enumerate(rows[:30]) if HEADERS[1] in row]
            if indexes:
                h = indexes[0]
                for n, row in enumerate(rows[h+1:], h+2):
                    yield 'CSV', n, rows[h], row, set()
                break
    elif path.suffix.lower() == '.xls':
        import xlrd
        # A BIFF formula has only its cached value in xlrd. Do not mistake it
        # for an engineer's literal correction (including formulas off-screen).
        from xlrd.compdoc import CompDoc
        import struct
        compound = CompDoc(path.read_bytes())
        stream = compound.get_named_stream('Workbook')
        if stream is None:
            stream = compound.get_named_stream('Book')
        position = 0
        while stream and position + 4 <= len(stream):
            kind, length = struct.unpack_from('<HH', stream, position)
            if kind in (0x0006, 0x0206, 0x0406):
                raise ValueError('Проверенный XLS содержит формулы. Сохраните проверяемые значения без формул или верните XLSX.')
            position += 4 + length
        wb = xlrd.open_workbook(path)
        try:
            for ws in wb.sheets():
                headers = [i for i in range(min(ws.nrows, 30)) if HEADERS[1] in ws.row_values(i)]
                if headers:
                    h = headers[0]
                    for n in range(h+1, ws.nrows):
                        yield ws.name, n+1, list(map(text_value, ws.row_values(h))), list(map(text_value, ws.row_values(n))), set()
        finally:
            wb.release_resources()
    else:
        raise ValueError('Верните Excel или CSV, обработанный этой программой.')


def import_review(path, store, progress=None, cancel=None):
    # Collect before committing: duplicate IDs are ambiguous even if just one
    # of their copies was edited. Never settle by first/last workbook order.
    staged, issues, untouched = [], [], 0
    store.sync()
    sessions = set()
    for sheet, num, header, values, formulas in read_return(path):
        if cancel and cancel.is_set():
            raise InterruptedError('Импорт отменён до сохранения решений')
        def get(label):
            indexes = [i for i, h in enumerate(header) if h == label]
            return values[indexes[-1]] if indexes and indexes[-1] < len(values) else ''
        try:
            _, cols = find_header([header])
        except ValueError:
            continue
        source = values[cols['name']] if cols['name'] < len(values) else ''
        if not source.strip():
            continue
        code = values[cols['code']] if 'code' in cols and cols['code'] < len(values) else ''
        factory = values[cols['factory']] if 'factory' in cols and cols['factory'] < len(values) else ''
        rid, sid = get('MTR_ROW_ID'), get('MTR_SESSION_ID')
        if sid and sid not in sessions:
            try:
                store.load_session(sid); sessions.add(sid)
            except Exception as exc:
                issues.append({'sheet': sheet, 'row': num, 'reason': 'Не удалось загрузить исходный сеанс: ' + str(exc)})
                continue
        row = store.row(rid) if rid else None
        if not rid:
            matches = store.match_rows(code, source, factory, sid or None)
            if len(matches) == 1:
                row = matches[0]
        reason = None
        important = {cols['name']} | ({cols['code']} if 'code' in cols else set()) | {i for i, h in enumerate(header) if h in (HEADERS[1], REVIEW_HEADERS[0], 'MTR_ROW_ID', 'MTR_SESSION_ID')}
        if formulas & important:
            reason = 'Формула или ошибка в проверяемых данных'
        elif row is None:
            reason = 'Нет однозначного соответствия исходной строке'
        elif code_key(code) != code_key(row['code']) or normalize(source) != normalize(row['source']) or normalize(factory) != normalize(row['factory']) or (sid and sid != row['session']):
            reason = 'Исходные данные или идентификатор изменены; решение не перенесено'
        staged.append(dict(row=row, rid=rid or (row['id'] if row else ''), sheet=sheet, number=num,
                           final=get(HEADERS[1]), action=get(REVIEW_HEADERS[0]).strip().casefold(), comment=get(REVIEW_HEADERS[1]), error=reason))
        if progress and len(staged) % 500 == 0:
            progress('Сопоставлено строк: ' + str(len(staged)))
    duplicates = Counter(x['rid'] for x in staged if x['rid'])
    accepted = repeated = 0
    for item in staged:
        error = item['error']
        if duplicates[item['rid']] > 1:
            error = 'Повтор MTR_ROW_ID; обе копии отправлены на ручную проверку'
        action = item['action']
        if action and action not in ACTIONS:
            error = 'Неизвестная отметка инженера: ' + action
        if error:
            issues.append({'sheet': item['sheet'], 'row': item['number'], 'reason': error}); continue
        row = item['row']; final = item['final']
        if action == 'оставить как в исходном':
            final = row['source']
        if not action and final == row['automatic']:
            untouched += 1; continue
        if not final.strip():
            issues.append({'sheet': item['sheet'], 'row': item['number'], 'reason': 'Пустое наименование не принято'}); continue
        if action == 'правильно' and final != row['automatic']:
            action = 'исправить'
        event = store.decide(dict(row, provenance=dict(row.get('provenance', {}), comment=item['comment'])), final, ACTIONS.get(action, 'Исправить'))
        accepted += event is not None; repeated += event is None
    if not staged and not issues:
        raise ValueError('В файле не найдены столбцы результата обезличивания.')
    status = store.sync()
    return {'accepted': accepted, 'repeated': repeated, 'untouched': untouched, 'issues': issues, 'sync': status}


def export_knowledge(path, snapshot):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'Знания'
    ws.append(['Ключ', 'Вид', 'Состояние', 'Решение', 'Независимых инженеров', 'Противоречащих', 'Область', 'Пример'])
    history = wb.create_sheet('История')
    history.append(['Ключ', 'ID', 'Дата', 'Автор', 'Действие', 'Значение', 'Исходное', 'Автоматическое', 'Удалено', 'Восстановлено', 'Основание'])
    for e in snapshot.entries.values():
        sample = e['sample']
        ws.append([e['key'], e['kind'], e['status'], e['value'], e['confirmations'], e['opposition'], str(sample.get('scope', {})), sample.get('source', sample.get('example', ''))])
        for h in e['history']:
            history.append([e['key'], h['id'], h['timestamp'], h['user'], h.get('action', h['kind']), h['value'], h.get('source', h.get('example', '')), h.get('automatic', ''), str(h.get('removed', [])), str(h.get('restored', [])), str(h.get('provenance', h.get('reason', '')))])
    for sheet in wb:
        sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for col in sheet[1]:
            sheet.column_dimensions[col.column_letter].width = 28
    wb.save(path); wb.close()
