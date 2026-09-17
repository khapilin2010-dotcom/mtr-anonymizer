"""Controlled Excel export/import for knowledge maintenance.

The exported workbook is intentionally not the database itself. Engineers edit
only explicit management columns; import validates the current key/value before
creating a new immutable event. Audit history is preserved.
"""
from pathlib import Path
import uuid

from excel.knowledge import digest

STATUS_NAMES = {'ACTIVE': 'ДЕЙСТВУЕТ', 'TRUSTED': 'ДЕЙСТВУЕТ',
                'DISABLED': 'ОТКЛЮЧЕНО', 'DISPUTED': 'ТРЕБУЕТ ПРОВЕРКИ', 'CANDIDATE': 'НА ПРОВЕРКЕ'}
RULE_NAMES = {'KEEP': 'ОСТАВИТЬ', 'DELETE': 'УДАЛИТЬ'}

def revision(entry):
    return digest([entry['status'], entry['value'], sorted(e['id'] for e in entry['history'])])

def display_value(entry):
    return RULE_NAMES.get(entry['value'], entry['value']) if entry['kind'] == 'rule' else entry['value']


HEADERS = [
    'Ключ', 'Тип', 'Статус', 'Текущее решение', 'Новое решение', 'Действие',
    'Независимых инженеров', 'Область', 'Пример', 'Причина изменения',
    'Код Автодокс', 'Исходное наименование', 'Завод', 'Тип решения',
    'Дата решения', 'Инженер', 'Версия записи'
]
ACTION_ALIASES = {
    'CHANGE': 'CHANGE', 'ИЗМЕНИТЬ': 'CHANGE',
    'DISABLE': 'DISABLE', 'ОТКЛЮЧИТЬ': 'DISABLE',
    'ENABLE': 'ENABLE', 'ВКЛЮЧИТЬ': 'ENABLE',
}


def export_editable(path, store):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    store.sync()
    snapshot = store.snapshot()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Управление'
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1767A6')
        cell.alignment = Alignment(wrap_text=True)
    type_names = {'case': 'Точное решение', 'rule': 'Общее правило', 'control': 'Настройка'}
    for key, entry in sorted(snapshot.entries.items(), key=lambda item: (item[1]['kind'] != 'case', item[1].get('sample', {}).get('code', ''), item[0])):
        sample = entry.get('sample', {})
        scope = sample.get('scope', {})
        example = sample.get('source', sample.get('example', ''))
        latest = max(entry.get('live') or entry['history'], key=lambda event: (event['timestamp'], event['id']))
        action = {'Правильно': 'Оставить результат программы', 'Исправить': 'Исправить результат'}.get(latest.get('action', ''), latest.get('action', ''))
        ws.append([key, 'Сохранение серии' if sample.get('series_rule') else type_names.get(entry.get('kind', ''), entry.get('kind', '')),
                   STATUS_NAMES.get(entry.get('status'), 'ТРЕБУЕТ ПРОВЕРКИ'), display_value(entry), '', '',
                   entry.get('confirmations', 0), str(scope), example, '',
                   sample.get('code', ''), example, sample.get('factory', scope.get('factory', '')),
                   action or 'Общее правило', latest.get('timestamp', ''),
                   ', '.join(sorted({e['user'] for e in entry.get('live', [])})) or latest.get('user', ''), revision(entry)])
    validation = DataValidation(
        type='list', formula1='"ИЗМЕНИТЬ,ОТКЛЮЧИТЬ,ВКЛЮЧИТЬ"', allow_blank=True)
    validation.errorTitle = 'Действие'
    validation.error = 'Выберите действие из списка.'
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    if ws.max_row >= 2:
        validation.add(f'F2:F{ws.max_row}')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    widths = [44, 18, 23, 55, 55, 16, 20, 38, 70, 35, 20, 70, 28, 30, 28, 28, 24]
    for index, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(index)].width = width

    for letter in ('A', 'H', 'I', 'Q'):
        ws.column_dimensions[letter].hidden = True
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    note = wb.create_sheet('Инструкция')
    note.append(['Как исправить базу решений'])
    note.append(['1. Найдите нужную строку по примеру, коду/наименованию или текущему решению.'])
    note.append(['2. Чтобы заменить решение: впишите правильный текст в «Новое решение», выберите «ИЗМЕНИТЬ» и укажите причину.'])
    note.append(['3. Если ошибочное правило вообще не должно работать: выберите «ОТКЛЮЧИТЬ» и укажите причину.'])
    note.append(['4. Чтобы вернуть ранее отключённую запись: выберите «ВКЛЮЧИТЬ» и укажите причину.'])
    note.append(['5. Не меняйте «Ключ» и «Текущее решение». Строки без действия при загрузке игнорируются.'])
    note.append(['6. Для общего правила новое решение при ИЗМЕНИТЬ должно быть ОСТАВИТЬ или УДАЛИТЬ.'])
    note.append(['7. История не стирается: исправление записывается новым событием, старое остаётся в аудите.'])
    note.column_dimensions['A'].width = 120

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def preview_editable(path, store):
    from openpyxl import load_workbook
    store.sync()
    snapshot = store.snapshot()
    wb = load_workbook(path, read_only=True, data_only=False)
    issues = []
    changes = []
    base = None
    try:
        if 'Управление' not in wb.sheetnames:
            raise ValueError('Нет листа «Управление». Используйте выгрузку этой программы.')
        ws = wb['Управление']
        header = [str(c.value or '').strip() for c in next(ws.iter_rows())]
        missing = [name for name in HEADERS[:10] if name not in header]
        if missing:
            raise ValueError('Не хватает столбцов: ' + ', '.join(missing))
        pos = {name: header.index(name) for name in HEADERS if name in header}
        seen = set()
        for row_no, cells in enumerate(ws.iter_rows(min_row=2), 2):
            values = [c.value for c in cells]
            raw_action = str(values[pos['Действие']] or '').strip().upper()
            if not raw_action:
                continue
            action = ACTION_ALIASES.get(raw_action)
            key = str(values[pos['Ключ']] or '').strip()
            current_file = str(values[pos['Текущее решение']] or '')
            new_value = str(values[pos['Новое решение']] or '')
            reason = str(values[pos['Причина изменения']] or '').strip()
            if not action:
                issues.append((row_no, 'Неизвестное действие: ' + raw_action)); continue
            entry = snapshot.entries.get(key)
            if not entry:
                issues.append((row_no, 'Ключ отсутствует в текущей базе')); continue
            if current_file not in (str(entry.get('value', '')), display_value(entry)):
                issues.append((row_no, 'База изменилась после выгрузки; строка устарела')); continue
            if 'Версия записи' in pos and str(values[pos['Версия записи']] or '') != revision(entry):
                issues.append((row_no, 'База изменилась после выгрузки; строка устарела')); continue
            if key in seen:
                issues.append((row_no, 'Повторное действие для того же решения')); continue
            seen.add(key)
            if any(cells[pos[name]].data_type == 'f' for name in ('Новое решение', 'Действие', 'Причина изменения')):
                issues.append((row_no, 'В полях изменения должны быть значения, а не формулы')); continue
            if not reason:
                issues.append((row_no, 'Не указана причина изменения')); continue
            if action == 'CHANGE':
                if not new_value.strip():
                    issues.append((row_no, 'Для «ИЗМЕНИТЬ» заполните новое решение')); continue
                if entry.get('kind') == 'rule':
                    new_value = {'ОСТАВИТЬ': 'KEEP', 'УДАЛИТЬ': 'DELETE'}.get(new_value.strip().upper(), new_value.strip().upper())
                if entry.get('kind') == 'rule' and new_value not in ('KEEP', 'DELETE'):
                    issues.append((row_no, 'Для общего правила допустимы только ОСТАВИТЬ или УДАЛИТЬ')); continue
            if action == 'CHANGE' and entry.get('kind') == 'case':
                from excel.expert_engine import protection_losses
                from excel.excel_engine import Anonymizer
                sample = entry['sample']
                base = base or Anonymizer()
                losses = protection_losses(sample.get('source', ''), new_value, base, sample.get('code', ''), sample.get('factory', ''))
                if losses:
                    issues.append((row_no, 'Потеряны защищённые признаки: ' + ', '.join(losses))); continue
            changes.append(dict(row=row_no, key=key, action=action,
                                new_value=new_value, reason=reason, entry=entry))
    finally:
        wb.close()
    return dict(ready=len(changes), issues=issues, changes=changes)


def commit_editable(path, store):
    plan = preview_editable(path, store)
    applied = 0
    for item in plan['changes']:
        store.sync(auto_backup=False)
        entry = store.snapshot().entries.get(item['key'])
        if not entry or revision(entry) != revision(item['entry']):
            plan['issues'].append((item['row'], 'База изменилась; выгрузите её заново'))
            continue
        action = item['action']
        reason = item['reason']
        if action in ('DISABLE', 'ENABLE'):
            store.control(item['key'], action == 'DISABLE', reason)
            applied += 1
            continue
        sample = entry.get('sample', {})
        if entry.get('kind') == 'case':
            row = dict(id=sample.get('row_id') or str(uuid.uuid4()),
                       session=sample.get('session') or ('knowledge-edit:' + uuid.uuid4().hex),
                       code=sample.get('code', ''), source=sample.get('source', ''),
                       factory=sample.get('factory', ''),
                       automatic=sample.get('automatic', sample.get('source', '')),
                       classification=sample.get('classification'), explicit_feedback=True,
                       provenance={'source': 'Редактирование базы через Excel', 'reason': reason})
            store.decide(row, item['new_value'], 'Редактирование базы Excel')
            applied += 1
        elif entry.get('kind') == 'rule':
            store.administrate(item['key'], reason, decision_override=item['new_value'],
                               requested_status='ACTIVE')
            applied += 1
    store.sync()
    return dict(applied=applied, issues=plan['issues'])
