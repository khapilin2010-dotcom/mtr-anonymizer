# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import os
import re
import sys
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

APP_TITLE = "Проверка опросных листов — заводы"
APP_VERSION = "0.1 RC1"

FACTORY_HEADERS = {
    "завод", "изготовитель", "производитель", "поставщик",
    "завод/изготовитель/поставщик", "завод изготовитель", "завод-изготовитель",
    "наименование завода", "наименование производителя"
}
LEGAL_RE = re.compile(r"(?iu)(?<![\w])(?:ООО|АО|ПАО|ОАО|ЗАО|НПО|НПП|ФГУП|ГУП|ИП|LLC|LTD\.?|LIMITED|INC\.?|CORP\.?|GMBH|AG|S\.?A\.?|B\.?V\.?)")
INN_RE = re.compile(r"(?iu)\bИНН\s*[:№#-]?\s*\d{10,12}\b")
KPP_RE = re.compile(r"(?iu)\bКПП\s*[:№#-]?\s*\d{9}\b")
ROLE_RE = re.compile(r"(?iu)\b(?:производитель|изготовитель|поставщик|завод[- ]?изготовитель)\s*[:\-–—]")
GENERIC = {"завод","производитель","изготовитель","поставщик","россия","рф","москва","компания","предприятие","торговый дом","ооо","ао","пао","оао","зао","нпо","нпп"}


def exe_dir() -> Path:
    return Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent


def norm(s: str) -> str:
    s = str(s or "").lower().replace("ё", "е")
    s = re.sub(r"[^0-9a-zа-я]+", " ", s, flags=re.I)
    return re.sub(r"\s+", " ", s).strip()


def norm_header(v) -> str:
    return re.sub(r"[\s_]+", " ", str(v or "").strip().lower().replace("ё", "е"))


def aliases(factory: str) -> List[str]:
    raw = re.sub(r"\s+", " ", str(factory or "")).strip(" ,;")
    if not raw:
        return []
    cut = re.split(r"(?iu)\b(?:ИНН|КПП|тел(?:ефон)?|e-?mail)\b|https?://|www\.|@", raw, maxsplit=1)[0].strip(" ,;.-")
    cand = {raw, cut, cut.split(",", 1)[0]}
    for p in (r"«([^»]{3,100})»", r'"([^"\n]{3,100})"', r"“([^”]{3,100})”"):
        cand.update(m.group(1).strip() for m in re.finditer(p, raw))
    stripped = re.sub(r"(?iu)^\s*(?:ООО|АО|ПАО|ОАО|ЗАО|НПО|НПП|ФГУП|ГУП|ИП|LLC|LTD\.?|LIMITED|INC\.?|CORP\.?|GMBH|AG|S\.?A\.?|B\.?V\.?)\s*", "", cut).strip(" \"'«».,;-")
    cand.add(stripped)
    cand.add(stripped.split(",", 1)[0])
    out = []
    for x in cand:
        n = norm(x)
        if n and n not in GENERIC and len(n) >= 4 and sum(c.isalpha() for c in n) >= 3:
            if n not in out:
                out.append(n)
    return sorted(out, key=len, reverse=True)


class Matcher:
    def __init__(self):
        self.by_first: Dict[str, List[tuple[str,str]]] = {}

    def add(self, alias: str, factory: str):
        if alias:
            self.by_first.setdefault(alias[0], []).append((alias, factory))

    def finish(self):
        for k in self.by_first:
            self.by_first[k].sort(key=lambda x: len(x[0]), reverse=True)

    def find(self, text: str):
        n = norm(text)
        found = []
        seen = set()
        for first, items in self.by_first.items():
            pos = n.find(first)
            while pos >= 0:
                if pos == 0 or not n[pos-1].isalnum():
                    for a, f in items:
                        if n.startswith(a, pos):
                            end = pos + len(a)
                            if end == len(n) or not n[end].isalnum():
                                key = (a, f)
                                if key not in seen:
                                    seen.add(key)
                                    found.append((a, f))
                                break
                pos = n.find(first, pos + 1)
        return found


class FactoryDB:
    def __init__(self):
        self.factories = []
        self.matcher = Matcher()
        self.source = None

    def load(self, path: Path):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        vals = []
        try:
            for ws in wb.worksheets:
                col = row0 = None
                for r in range(1, min(ws.max_row or 1, 40) + 1):
                    for c in range(1, min(ws.max_column or 1, 100) + 1):
                        if norm_header(ws.cell(r, c).value) in FACTORY_HEADERS:
                            row0, col = r, c
                            break
                    if col:
                        break
                if not col:
                    continue
                for r in range(row0 + 1, (ws.max_row or row0) + 1):
                    v = re.sub(r"\s+", " ", str(ws.cell(r, col).value or "")).strip()
                    if v:
                        vals.append(v)
        finally:
            wb.close()
        if not vals:
            raise ValueError("В Excel не найден заполненный столбец Завод/Изготовитель/Производитель.")
        self.factories = list(dict.fromkeys(vals))
        self.matcher = Matcher()
        for f in self.factories:
            for a in aliases(f):
                self.matcher.add(a, f)
        self.matcher.finish()
        self.source = path


@dataclass
class Hit:
    path: str
    file: str
    page: int
    level: str
    found: str
    kind: str
    factory: str
    context: str
    decision: str = ""


def context_for(text: str, needle: str) -> str:
    p = norm(text).find(needle)
    if p < 0:
        return re.sub(r"\s+", " ", text)[:240]
    # Exact char mapping is unnecessary for the review pane; find the first token in raw text.
    token = needle.split()[0]
    m = re.search(re.escape(token), text, re.I)
    if not m:
        return re.sub(r"\s+", " ", text)[:240]
    a, b = max(0, m.start()-100), min(len(text), m.end()+140)
    return re.sub(r"\s+", " ", text[a:b]).strip()


def scan_pdf(path: Path, db: FactoryDB, progress=None):
    import fitz
    doc = fitz.open(path)
    hits = []
    no_text = 0
    pages = doc.page_count
    try:
        for i in range(pages):
            text = doc[i].get_text("text") or ""
            if not text.strip():
                no_text += 1
                if progress: progress(i+1, pages)
                continue
            red_factories = set()
            for a, f in db.matcher.find(text):
                red_factories.add(f)
                hits.append(Hit(str(path), path.name, i+1, "RED", a, "База заводов", f, context_for(text, a)))
            yellow = []
            for rx, kind in ((INN_RE,"ИНН"),(KPP_RE,"КПП"),(ROLE_RE,"Роль организации")):
                yellow += [(m.group(0), kind) for m in rx.finditer(text)]
            for m in LEGAL_RE.finditer(text):
                tail = text[m.end():m.end()+90]
                tail = re.split(r"[\n;]|\s{3,}", tail, maxsplit=1)[0].strip()
                yellow.append(((m.group(0) + (" " + tail if tail else ""))[:100], "Организация"))
            seen = set()
            for raw, kind in yellow:
                nr = norm(raw)
                if not nr or nr in seen:
                    continue
                seen.add(nr)
                # If this yellow fragment already contains a confirmed factory alias, red is enough.
                if db.matcher.find(raw):
                    continue
                hits.append(Hit(str(path), path.name, i+1, "YELLOW", raw, kind, "", context_for(text, norm(raw))))
            if progress: progress(i+1, pages)
    finally:
        doc.close()
    return hits, no_text, pages


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_TITLE} {APP_VERSION}")
        self.geometry("1420x820")
        self.minsize(1050, 650)
        self.db = FactoryDB()
        self.pdfs: List[Path] = []
        self.hits: List[Hit] = []
        self.running = False
        self.base_var = tk.StringVar(value="База заводов: не загружена")
        self.status_var = tk.StringVar(value="Готово")
        self.build_ui()
        self.after(250, self.auto_base)

    def build_ui(self):
        top = ttk.Frame(self, padding=8); top.pack(fill="x")
        ttk.Button(top,text="1. База заводов",command=self.choose_base).pack(side="left",padx=3)
        ttk.Button(top,text="2. Добавить PDF",command=self.add_pdf).pack(side="left",padx=3)
        ttk.Button(top,text="Добавить папку",command=self.add_folder).pack(side="left",padx=3)
        ttk.Button(top,text="3. Проверить",command=self.start).pack(side="left",padx=12)
        ttk.Button(top,text="Экспорт отчёта",command=self.export).pack(side="left",padx=3)
        ttk.Button(top,text="Очистить",command=self.clear).pack(side="left",padx=3)
        ttk.Label(self,textvariable=self.base_var,padding=(10,0)).pack(fill="x")
        ttk.Label(self,textvariable=self.status_var,padding=(10,4)).pack(fill="x")
        pane = ttk.Panedwindow(self,orient="horizontal"); pane.pack(fill="both",expand=True,padx=8,pady=8)
        left,right=ttk.Frame(pane),ttk.Frame(pane); pane.add(left,weight=3); pane.add(right,weight=2)
        cols=("level","file","page","found","kind","decision")
        self.tree=ttk.Treeview(left,columns=cols,show="headings",selectmode="browse")
        heads={"level":"Статус","file":"Файл","page":"Стр.","found":"Найдено","kind":"Основание","decision":"Решение"}
        widths={"level":80,"file":190,"page":55,"found":310,"kind":135,"decision":100}
        for c in cols: self.tree.heading(c,text=heads[c]); self.tree.column(c,width=widths[c],stretch=c in {"file","found"})
        self.tree.tag_configure("red",background="#F4CCCC"); self.tree.tag_configure("yellow",background="#FFF2CC"); self.tree.tag_configure("done",background="#D9EAD3")
        sy=ttk.Scrollbar(left,orient="vertical",command=self.tree.yview); self.tree.configure(yscrollcommand=sy.set)
        self.tree.pack(side="left",fill="both",expand=True); sy.pack(side="right",fill="y")
        self.tree.bind("<<TreeviewSelect>>",self.select)
        bar=ttk.Frame(right); bar.pack(fill="x",pady=(0,6))
        ttk.Button(bar,text="Удалить",command=lambda:self.decision("УДАЛИТЬ")).pack(side="left",padx=3)
        ttk.Button(bar,text="Оставить",command=lambda:self.decision("ОСТАВИТЬ")).pack(side="left",padx=3)
        ttk.Button(bar,text="Сбросить",command=lambda:self.decision("")).pack(side="left",padx=3)
        ttk.Button(bar,text="Открыть PDF",command=self.open_pdf).pack(side="left",padx=12)
        ttk.Label(bar,text="PDF не изменяется").pack(side="right",padx=4)
        self.info=tk.Text(right,wrap="word",font=("Segoe UI",11)); self.info.pack(fill="both",expand=True)
        self.info.configure(state="disabled")
        self.progress=ttk.Progressbar(self,mode="determinate"); self.progress.pack(fill="x",padx=8,pady=(0,8))

    def auto_base(self):
        roots=[exe_dir(),exe_dir().parent]
        cand=[]
        for r in roots:
            if r.exists(): cand += list(r.glob("*База*обезличив*.xlsx"))+list(r.glob("*BASELINE*.xlsx"))
        if cand: self.load_base(sorted(set(cand),key=lambda p:("BASELINE" not in p.name.upper(),len(p.name)))[0],True)

    def choose_base(self):
        p=filedialog.askopenfilename(title="Excel с базой заводов",filetypes=[("Excel","*.xlsx *.xlsm")])
        if p: self.load_base(Path(p),False)

    def load_base(self,p:Path,quiet=False):
        try:
            self.status_var.set("Загружаю базу заводов…"); self.update_idletasks(); self.db.load(p)
            self.base_var.set(f"База: {p.name} | уникальных заводов: {len(self.db.factories)}")
            self.status_var.set("База загружена")
        except Exception as e:
            self.base_var.set("База заводов: не загружена"); self.status_var.set(str(e))
            if not quiet: messagebox.showerror(APP_TITLE,str(e))

    def add_pdf(self):
        self.add_paths(filedialog.askopenfilenames(title="PDF",filetypes=[("PDF","*.pdf")]))

    def add_folder(self):
        p=filedialog.askdirectory(title="Папка со сборником PDF")
        if p: self.add_paths([str(x) for x in Path(p).rglob("*.pdf")])

    def add_paths(self,paths):
        have={str(p.resolve()).lower() for p in self.pdfs}
        for x in paths:
            p=Path(x); k=str(p.resolve()).lower()
            if k not in have: self.pdfs.append(p); have.add(k)
        self.status_var.set(f"PDF в очереди: {len(self.pdfs)}")

    def clear(self):
        if self.running: return
        self.pdfs.clear(); self.hits.clear(); self.progress["value"]=0
        for x in self.tree.get_children(): self.tree.delete(x)
        self.set_info(""); self.status_var.set("Очищено")

    def start(self):
        if self.running: return
        if not self.db.factories: messagebox.showwarning(APP_TITLE,"Сначала выберите базу заводов."); return
        if not self.pdfs: messagebox.showwarning(APP_TITLE,"Добавьте PDF или папку."); return
        self.running=True; self.hits=[]
        for x in self.tree.get_children(): self.tree.delete(x)
        threading.Thread(target=self.worker,daemon=True).start()

    def worker(self):
        all_hits=[]; no_text=0; errors=[]; total=len(self.pdfs)
        for fi,p in enumerate(self.pdfs,1):
            try:
                def prog(pg,pages):
                    val=((fi-1)+pg/max(1,pages))/max(1,total)*100
                    self.after(0,lambda v=val,t=f"{p.name}: страница {pg} из {pages}":self.progress_ui(v,t))
                h,n,_=scan_pdf(p,self.db,prog); all_hits+=h; no_text+=n
            except Exception as e: errors.append(f"{p.name}: {e}")
        self.after(0,lambda:self.done(all_hits,no_text,errors))

    def progress_ui(self,v,t): self.progress["value"]=v; self.status_var.set(t)

    def done(self,hits,no_text,errors):
        self.running=False; self.hits=hits
        for i,h in enumerate(hits):
            self.tree.insert("","end",iid=str(i),values=("КРАСНЫЙ" if h.level=="RED" else "ЖЁЛТЫЙ",h.file,h.page,h.found,h.kind,h.decision),tags=(("red" if h.level=="RED" else "yellow"),))
        red=sum(h.level=="RED" for h in hits); yellow=len(hits)-red
        s=f"Готово: красных {red}, жёлтых {yellow}, файлов {len(self.pdfs)}"
        if no_text: s+=f". Страниц без текстового слоя: {no_text}"
        if errors: s+=f". Ошибок: {len(errors)}"
        self.status_var.set(s); self.progress["value"]=100
        if errors: messagebox.showwarning(APP_TITLE,"Не обработано:\n\n"+"\n".join(errors[:10]))
        if self.tree.get_children(): self.tree.selection_set(self.tree.get_children()[0]); self.select()

    def select(self,_=None):
        sel=self.tree.selection()
        if not sel: return
        h=self.hits[int(sel[0])]
        txt=f"Файл: {h.file}\nСтраница: {h.page}\nСтатус: {'КРАСНЫЙ' if h.level=='RED' else 'ЖЁЛТЫЙ'}\nНайдено: {h.found}\nОснование: {h.kind}"
        if h.factory: txt+=f"\n\nЗавод из базы:\n{h.factory}"
        txt+=f"\n\nКонтекст страницы:\n{h.context}\n\nРешение инженера: {h.decision or 'не принято'}"
        self.set_info(txt)

    def set_info(self,t):
        self.info.configure(state="normal"); self.info.delete("1.0","end"); self.info.insert("1.0",t); self.info.configure(state="disabled")

    def decision(self,d):
        sel=self.tree.selection()
        if not sel: return
        idx=int(sel[0]); h=self.hits[idx]; h.decision=d
        v=list(self.tree.item(sel[0],"values")); v[-1]=d
        self.tree.item(sel[0],values=v,tags=(("done" if d else ("red" if h.level=="RED" else "yellow")),)); self.select()

    def open_pdf(self):
        sel=self.tree.selection()
        if not sel: return
        h=self.hits[int(sel[0])]
        try:
            if os.name=="nt": os.startfile(h.path)
            elif sys.platform=="darwin": os.system(f'open "{h.path}"')
            else: os.system(f'xdg-open "{h.path}" >/dev/null 2>&1 &')
        except Exception as e: messagebox.showerror(APP_TITLE,str(e))

    def export(self):
        if not self.hits: messagebox.showinfo(APP_TITLE,"Нет результатов."); return
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",initialfile="Проверка_опросных_листов.xlsx",filetypes=[("Excel","*.xlsx"),("CSV","*.csv")])
        if not p: return
        try:
            if Path(p).suffix.lower()==".csv":
                with open(p,"w",encoding="utf-8-sig",newline="") as f:
                    w=csv.writer(f,delimiter=";"); w.writerow(["Статус","Файл","Страница","Найдено","Основание","Завод из базы","Решение","Контекст"])
                    for h in self.hits: w.writerow([h.level,h.file,h.page,h.found,h.kind,h.factory,h.decision,h.context])
            else:
                from openpyxl import Workbook
                from openpyxl.styles import PatternFill,Font
                wb=Workbook(); ws=wb.active; ws.title="Проверка"; ws.append(["Статус","Файл","Страница","Найдено","Основание","Завод из базы","Решение","Контекст"])
                for c in ws[1]: c.font=Font(bold=True)
                fills={"RED":PatternFill("solid",fgColor="F4CCCC"),"YELLOW":PatternFill("solid",fgColor="FFF2CC"),"DONE":PatternFill("solid",fgColor="D9EAD3")}
                for h in self.hits:
                    ws.append(["КРАСНЫЙ" if h.level=="RED" else "ЖЁЛТЫЙ",h.file,h.page,h.found,h.kind,h.factory,h.decision,h.context]); fill=fills["DONE"] if h.decision else fills[h.level]
                    for c in ws[ws.max_row]: c.fill=fill
                for col,w in {"A":12,"B":28,"C":10,"D":35,"E":20,"F":48,"G":14,"H":75}.items(): ws.column_dimensions[col].width=w
                ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; wb.save(p)
            self.status_var.set(f"Отчёт сохранён: {Path(p).name}")
        except Exception as e: messagebox.showerror(APP_TITLE,str(e))


if __name__=="__main__":
    App().mainloop()
