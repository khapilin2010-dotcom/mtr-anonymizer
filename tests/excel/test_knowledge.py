"""Synthetic acceptance tests for shared decisions, review return and recovery."""
import concurrent.futures
import gzip
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

from excel.knowledge import KnowledgeStore, Snapshot, case_key, code_key, canonical, digest, restore_backup
from excel.expert_engine import ExpertAnonymizer
from excel.excel_engine import Anonymizer
from excel.file_io import process_file, HEADERS
from excel.review_io import import_review, REVIEW_HEADERS, export_knowledge

class KnowledgeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.base = Anonymizer()
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.share = self.root / 'Общая папка'; self.share.mkdir()
        self.a = KnowledgeStore(self.root / 'alice', self.share, 'DOMAIN\\alice')
        self.b = KnowledgeStore(self.root / 'bob', self.share, 'DOMAIN\\bob')
        self.c = KnowledgeStore(self.root / 'carol', self.share, 'DOMAIN\\carol')
        self.row = dict(id='synthetic-row', session='synthetic-session', code='SYNTH-001', source='Клапан XYZ-500 DN50', factory='Тестовый завод', automatic='Клапан XYZ-500 DN50', provenance={})
    def tearDown(self): self.tmp.cleanup()
    def result(self, store=None, source=None):
        return ExpertAnonymizer((store or self.a).snapshot(), self.base).anonymize(source or self.row['source'], self.row['code'], self.row['factory'])
    def entry(self, store=None): return (store or self.a).snapshot().entries[case_key(self.row['code'], self.row['source'], self.row['factory'])]
    def test_exact_case_immediate_and_immutable_batch(self):
        frozen = ExpertAnonymizer(self.a.snapshot(), self.base)
        self.a.decide(self.row, 'Клапан DN50')
        self.assertEqual(self.result()['text'], 'Клапан DN50'); self.assertEqual(self.result()['status'], 'ЗЕЛЁНЫЙ')
        self.assertEqual(frozen.anonymize(self.row['source'], self.row['code'], self.row['factory'])['text'], self.row['source'])
    def test_shared_three_independent_confirmations(self):
        for store in (self.a,self.b,self.c): store.decide(self.row, 'Клапан DN50'); store.sync()
        self.a.sync(); self.assertEqual(self.entry()['status'], 'TRUSTED'); self.assertEqual(self.entry()['confirmations'], 3)
    def test_repeat_same_user_never_inflates(self):
        for _ in range(30): self.a.decide(self.row, 'Клапан DN50')
        self.assertEqual(len(self.a.events()), 1); self.assertEqual(self.entry()['confirmations'], 1)
    def test_opposition_is_red_without_last_writer_wins(self):
        self.a.decide(self.row, 'Клапан DN50'); self.b.decide(self.row, self.row['source'])
        self.a.sync(); self.b.sync(); self.a.sync()
        self.assertEqual(self.entry()['status'], 'DISPUTED'); self.assertEqual(self.result()['status'], 'КРАСНЫЙ')
        self.assertEqual(self.result()['text'], self.row['source'])
    def test_own_correction_supersedes_only_own_vote(self):
        self.a.decide(self.row, 'Клапан DN50'); self.b.decide(self.row, self.row['source'])
        self.a.sync(); self.b.sync(); self.a.sync(); self.a.decide(self.row, self.row['source'])
        self.assertEqual(self.entry()['confirmations'], 2); self.assertEqual(self.entry()['status'], 'ACTIVE')
        self.assertEqual(len(self.entry()['history']), 3)
    def test_reducer_independent_of_order(self):
        self.a.decide(self.row, 'Клапан DN50'); self.a.decide(self.row, self.row['source'])
        events = self.a.events(); key = case_key(self.row['code'], self.row['source'], self.row['factory'])
        self.assertEqual(Snapshot(events).entries[key]['value'], Snapshot(list(reversed(events))).entries[key]['value'])
        self.assertEqual(Snapshot(events).version, Snapshot(list(reversed(events))).version)
    def test_offline_outbox_recovery(self):
        gone = self.root / 'disconnected'; self.share.rename(gone)
        self.a.decide(self.row, 'Клапан DN50'); self.assertFalse(self.a.sync()['online']); self.assertEqual(self.a.pending(), 1)
        self.assertFalse(self.share.exists()); gone.rename(self.share)
        self.assertTrue(self.a.sync()['online']); self.b.sync()
        self.assertEqual(self.result(self.b)['text'], 'Клапан DN50'); self.assertEqual(self.a.pending(), 0)
    def test_no_shared_sqlite(self):
        self.a.decide(self.row, 'Клапан DN50'); self.a.sync()
        self.assertFalse(list(self.share.rglob('*.sqlite*')))
    def test_partial_and_corrupt_events_ignored(self):
        folder=self.share/'events'/'aa'; folder.mkdir(parents=True)
        (folder/'partial.tmp').write_text('{'); (folder/'corrupt.json').write_text('{')
        result=self.a.sync(); self.assertTrue(result['online']); self.assertEqual(len(result['errors']),1); self.assertFalse(self.a.events())
    def test_checksum_failure_not_applied(self):
        event=self.a.decide(self.row,'Клапан DN50'); self.a.sync()
        path=next((self.share/'events').glob('*/*.json')); payload=json.loads(path.read_text('utf-8')); payload['event']['value']='BROKEN'; path.write_text(canonical(payload),encoding='utf-8')
        result=self.b.sync(); self.assertTrue(result['errors']); self.assertFalse(self.b.events())
    def test_wrong_store_never_mixes_data(self):
        self.a.decide(self.row, 'Клапан DN50')
        path=self.share/'mtr-knowledge.json'; data=json.loads(path.read_text('utf-8')); data['store_id']='00000000-0000-0000-0000-000000000000'; path.write_text(canonical(data),encoding='utf-8')
        self.assertFalse(self.a.sync()['online']); self.assertEqual(self.a.pending(),1)
    def test_concurrent_writers_keep_every_event(self):
        def work(store):
            for i in range(8): store.decide(dict(self.row,code='CONCURRENT-'+str(i)), 'Клапан DN50')
            return store.sync()
        with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool: results=list(pool.map(work,[self.a,self.b,self.c]))
        self.assertTrue(all(x['online'] for x in results)); self.a.sync(); self.assertEqual(len(self.a.events()),24)
        self.assertTrue(all(e['confirmations']==3 for e in self.a.snapshot().entries.values()))
    def test_protection_conflict_requires_review(self):
        self.a.decide(self.row,'Клапан'); result=self.result()
        self.assertEqual(result['status'],'КРАСНЫЙ'); self.assertIn('DN50',result['text']); self.assertIn('DN50',result['reason'])
    def test_aero_preserved_after_expert_rule(self):
        row=dict(self.row, code='AERO-SYNTHETIC', source='Установка АЭРО ИКСИА СКВ BOX(S)-S-63 IP54', factory='ИНН 3257017280')
        self.a.decide(row,'Установка IP54'); result=ExpertAnonymizer(self.a.snapshot(),self.base).anonymize(row['source'],row['code'],row['factory'])
        self.assertEqual(result['status'],'КРАСНЫЙ'); self.assertIn('BOX(S)-S-63',result['text'])
    def test_code_alias_reuses_same_decision(self):
        self.assertEqual(code_key('631-336532'),code_key('336532'))
        self.assertNotEqual(code_key('1-100-10'),code_key('100'))
    def test_without_code_requires_exact_source_factory(self):
        row=dict(self.row,code=''); self.a.decide(row,'Клапан DN50'); az=ExpertAnonymizer(self.a.snapshot(),self.base)
        self.assertEqual(az.anonymize(row['source'],'',row['factory'])['text'],'Клапан DN50')
        self.assertIn('XYZ-500',az.anonymize(row['source'],'','Другой завод')['text'])
    def test_generic_single_vote_is_only_candidate(self):
        self.a.propose_rule(self.row,'XYZ-500','DELETE','Клапан'); self.assertEqual(self.result()['text'],self.row['source'])
        entry=next(iter(self.a.snapshot().entries.values())); self.assertEqual(entry['status'],'CANDIDATE')
    def active_rule(self, fragment='XYZ-500', action='DELETE', row=None):
        row=row or self.row
        for store in (self.a,self.b): store.propose_rule(row,fragment,action,'Клапан'); store.sync()
        self.a.sync()
    def test_generic_rule_requires_scope_and_two_people(self):
        self.active_rule(); az=ExpertAnonymizer(self.a.snapshot(),self.base)
        self.assertEqual(self.result()['text'],'Клапан DN50')
        self.assertIn('XYZ-500',az.anonymize(self.row['source'],'','Другой завод')['text'])
        self.assertIn('XYZ-500',az.anonymize('Датчик XYZ-500 DN50','',self.row['factory'])['text'])
    def test_keep_rule_restores_deleted_legacy_brand(self):
        row=dict(self.row,source='Клапан Унипол DN50'); self.active_rule('Унипол','KEEP',row)
        az=ExpertAnonymizer(self.a.snapshot(),self.base); result=az.anonymize(row['source'],'',row['factory'])
        self.assertIn('Унипол',result['text'])
    def test_delete_rule_cannot_remove_technical(self):
        self.active_rule('DN50'); result=self.result(); self.assertEqual(result['status'],'КРАСНЫЙ'); self.assertIn('DN50',result['text'])
    def test_role_component_does_not_delete_product(self):
        for store in (self.a,self.b): store.propose_rule(self.row,'XYZ-500','DELETE','Клапан','комплектующее'); store.sync()
        self.a.sync(); self.assertEqual(self.result()['text'],self.row['source'])
    def test_component_rule_stays_inside_explicit_component_clause(self):
        row=dict(self.row,source='Клапан XYZ-500 DN50; привод XYZ-500 Тестовый завод')
        for store in (self.a,self.b):
            store.propose_rule(row,'XYZ-500','DELETE','привод','комплектующее'); store.sync()
        self.a.sync()
        result=ExpertAnonymizer(self.a.snapshot(),self.base).anonymize(row['source'],'','Другой изготовитель изделия')
        self.assertTrue(result['text'].startswith('Клапан XYZ-500 DN50'))
        self.assertEqual(result['text'].count('XYZ-500'),1)
    def test_automatic_backups_follow_changed_knowledge(self):
        self.a.decide(self.row,'Клапан DN50'); self.a.sync()
        self.assertTrue(list((self.share/'backups').glob('*.zip')))
    def test_disabled_knowledge_stops_applying_with_history(self):
        self.a.decide(self.row,'Клапан DN50'); key=self.entry()['key']; self.a.control(key,True,'Исправление ошибки')
        self.assertEqual(self.entry()['status'],'DISABLED'); self.assertEqual(self.result()['text'],self.row['source'])
        self.a.control(key,False,'Проверено'); self.assertEqual(self.result()['text'],'Клапан DN50')
    def test_non_admin_cannot_disable(self):
        with self.assertRaises(PermissionError): self.b.control('code:001',True,'test')
    def test_backup_restore_and_rebuild_cache(self):
        self.a.decide(self.row,'Клапан DN50'); self.a.sync(); path=self.root/'backup.zip'; self.a.backup(path)
        folder=self.root/'restored'; restore_backup(path,folder); fresh=KnowledgeStore(self.root/'fresh',folder,'reader'); fresh.sync()
        self.assertEqual(self.result(fresh)['text'],'Клапан DN50')
        with self.assertRaises(ValueError): restore_backup(path,folder)
    def test_export_history_includes_authors_and_restore(self):
        self.a.decide(self.row,'Клапан DN50'); self.a.decide(self.row,self.row['source']); path=self.root/'history.xlsx'
        export_knowledge(path,self.a.snapshot()); wb=load_workbook(path)
        self.assertEqual(wb['История'].max_row,3); self.assertEqual(wb['История'].cell(2,4).value,self.a.user); wb.close()

class ReturnTests(KnowledgeTests):
    # Inherited tests intentionally excluded below by unittest loader guard.
    def create(self, rows=None):
        source=self.root/'input.xlsx'; wb=Workbook(); ws=wb.active; ws.title='МТР'
        ws.append(['Код Автодокс','Наименование','Производитель'])
        for row in rows or [[self.row['code'],self.row['source'],self.row['factory']]]: ws.append(row)
        wb.save(source); original=source.read_bytes()
        result,report=process_file(source,self.root,ExpertAnonymizer(self.a.snapshot(),self.base),knowledge_store=self.a)
        self.assertEqual(original,source.read_bytes()); self.a.sync(); return result,report
    def test_return_untouched_is_not_confirmation(self):
        path,_=self.create(); report=import_review(path,self.a); self.assertEqual(report['untouched'],1); self.assertFalse(self.a.events())
    def test_corrected_excel_reused_on_another_machine(self):
        path,report=self.create(); wb=load_workbook(path); wb.active.cell(2,5,'Клапан DN50'); wb.save(path); wb.close()
        report=import_review(path,self.b); self.assertEqual(report['accepted'],1); self.assertFalse(report['issues'])
        self.a.sync(); self.assertEqual(self.result()['text'],'Клапан DN50')
        repeated=import_review(path,self.b); self.assertEqual(repeated['accepted'],0); self.assertEqual(repeated['repeated'],1)
    def test_explicit_unchanged_confirmation(self):
        path,_=self.create(); wb=load_workbook(path); wb.active.cell(2,8,'Правильно'); wb.save(path); wb.close()
        report=import_review(path,self.a); self.assertEqual(report['accepted'],1); self.assertEqual(self.entry()['value'],self.row['source'])
    def test_sorting_preserves_row_identity(self):
        rows=[[self.row['code'],self.row['source'],self.row['factory']],['SECOND','Клапан OTHER-500 DN50',self.row['factory']]]
        path,_=self.create(rows); wb=load_workbook(path); ws=wb.active; a=[c.value for c in ws[2]]; b=[c.value for c in ws[3]]
        for num,values in [(2,b),(3,a)]:
            for col,value in enumerate(values,1): ws.cell(num,col,value)
        ws.cell(3,5,'Клапан DN50'); wb.save(path); wb.close()
        report=import_review(path,self.b); self.assertFalse(report['issues']); self.assertEqual(report['accepted'],1)
    def test_duplicate_ids_refuse_both_copies(self):
        path,_=self.create(); wb=load_workbook(path); ws=wb.active; ws.append([c.value for c in ws[2]]); ws.cell(2,5,'Клапан DN50'); wb.save(path); wb.close()
        report=import_review(path,self.a); self.assertEqual(report['accepted'],0); self.assertEqual(len(report['issues']),2)
    def test_missing_ids_unambiguous_fallback(self):
        path,_=self.create(); wb=load_workbook(path); ws=wb.active; ws.cell(2,12).value=None; wb.save(path); wb.close()
        report=import_review(path,self.a); self.assertFalse(report['issues']); self.assertEqual(report['untouched'],1)
    def test_changed_source_never_taught_to_wrong_code(self):
        path,_=self.create(); wb=load_workbook(path); ws=wb.active; ws.cell(2,2,'Другой ресурс'); ws.cell(2,8,'Правильно'); wb.save(path); wb.close()
        report=import_review(path,self.a); self.assertEqual(report['accepted'],0); self.assertEqual(len(report['issues']),1)
    def test_hidden_ids_and_existing_four_columns(self):
        path,_=self.create(); wb=load_workbook(path); ws=wb.active
        self.assertEqual([c.value for c in ws[1]][3:7],list(HEADERS)); self.assertEqual([c.value for c in ws[1]][7:],list(REVIEW_HEADERS))
        self.assertTrue(ws.column_dimensions['L'].hidden); self.assertTrue(ws.column_dimensions['M'].hidden); wb.close()
    def test_corrupt_baseline_not_loaded_partially(self):
        path,report=self.create(); sid=report['session']; baseline=self.share/'sessions'/(sid+'.jsonl.gz')
        content=gzip.decompress(baseline.read_bytes()).decode('utf-8').splitlines(); content[-1]=canonical({'footer':999,'sha256':'bad'})
        baseline.write_bytes(gzip.compress(('\n'.join(content)+'\n').encode('utf-8')))
        with self.assertRaises(ValueError): self.b.load_session(sid)
        self.assertFalse(self.b.session_rows(sid))
    def test_cancel_output_discards_baseline_and_partial_file(self):
        src=self.root/'cancel.xlsx'; wb=Workbook(); wb.active.append(['Код','Наименование']); wb.active.append(['1','Клапан DN50']); wb.save(src)
        cancel=threading.Event(); cancel.set()
        with self.assertRaises(InterruptedError): process_file(src,self.root,knowledge_store=self.a,cancel=cancel)
        with self.a.db() as db: self.assertEqual(db.execute('SELECT count(*) FROM sessions').fetchone()[0],0)
        self.assertEqual(list(self.root.glob('*обезличено*')),[])

# Do not repeat all base tests in the subclass.
for _name in list(KnowledgeTests.__dict__):
    if _name.startswith('test_') and _name not in ReturnTests.__dict__:
        setattr(ReturnTests,_name,None)

if __name__ == '__main__': unittest.main()
