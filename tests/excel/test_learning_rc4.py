import uuid
import pytest
from excel.simple_store import SimpleKnowledgeStore
from excel.operator_engine import OperatorExpertAnonymizer
from excel.excel_engine import Anonymizer
from excel.knowledge import case_key
from excel.semantics import features

@pytest.fixture
def store(tmp_path):
    return SimpleKnowledgeStore(tmp_path/'local', tmp_path/'shared', 'engineer')

def record(store, source, final=None, factory='', code='1', action='Оставить как в исходном'):
    family = 'ПКМ-ТСТ' if 'ПКМ-ТСТ' in source else 'МОС'
    automatic = source.replace(family, '')
    row=dict(id=str(uuid.uuid4()),session='s',source=source,automatic=automatic,
             factory=factory,code=code,classification=features(source,factory))
    store.decide(row,source if final is None else final,action)
    return row

class RemovesFamilies(Anonymizer):
    def anonymize(self, source, code='', factory='', **kwargs):
        # Use the real span subtraction and preservation logic with explicit deletions.
        from excel.series_learning import family_spans
        deletes=[(a,b,'test rule') for family in ('МОС','ПКМ-ТСТ') for a,b in family_spans(source,family,True)]
        return super().anonymize(source,code,factory,**dict(kwargs,expert_deletes=deletes+list(kwargs.get('expert_deletes',()))))

def result(store,source,factory='',code='NEW'):
    return OperatorExpertAnonymizer(store.snapshot(),base=RemovesFamilies(),auto_apply_confirmed=True).anonymize(source,code,factory)

@pytest.mark.parametrize('family,head',[('МОС','Модуль автоматизированной технологической обвязки скважин'),('ПКМ-ТСТ','Панель')])
def test_one_engineer_keep_applies_to_new_models(store,family,head):
    record(store,f'{head} {family}-1',code='1')
    record(store,f'{head} {family}-2',code='2')
    source=f'{head} {family}-3 IP66'
    got=result(store,source)
    assert got['text']==source
    assert got['knowledge']['series_kept']==[family]
    assert not got['knowledge'].get('exact_applied')


def test_no_global_keep_other_equipment_or_factory_or_prefix(store):
    record(store,'Модуль МОС-1',factory='Завод А')
    for source,factory in [('Насос МОС-2','Завод А'),('Модуль МОС-2','Завод Б'),('Модуль МОС-2',''),('Модуль МОСКВА-2','Завод А')]:
        assert not result(store,source,factory)['knowledge'].get('series_kept')


def test_disagreement_requires_review_and_exact_still_wins(store):
    record(store,'Модуль МОС-1',code='1')
    record(store,'Модуль МОС-2',final='Модуль',code='2',action='Правильно')
    got=result(store,'Модуль МОС-3')
    assert got['status']=='КРАСНЫЙ' and got['knowledge']['series_conflicts']
    assert result(store,'Модуль МОС-1',code='1')['text']=='Модуль МОС-1'


def test_disabling_source_or_learned_rule_stops_generalization(store):
    row=record(store,'Модуль МОС-1')
    key=next(k for k in store.snapshot().entries if k.startswith('series:'))
    store.control(key,True,'Ошибочное обобщение')
    assert not result(store,'Модуль МОС-2')['knowledge'].get('series_kept')
    store.control(key,False,'Проверено')
    assert result(store,'Модуль МОС-2')['knowledge']['series_kept']
    store.control(case_key(row['code'],row['source']),True,'Ошибочная строка')
    assert not result(store,'Модуль МОС-2')['knowledge'].get('series_kept')


def test_mos_drawing_and_model_are_kept_on_next_machine(store,tmp_path):
    record(store,'Модуль МОС-3/1 Чертеж МОС-29.М1');store.sync(auto_backup=False)
    other=SimpleKnowledgeStore(tmp_path/'other',store.shared,'other');other.sync(auto_backup=False)
    source='Модуль МОС-4/2 Чертеж МОС-30.М2'
    assert result(other,source)['text']==source


def test_same_batch_refresh_writes_new_result_without_touching_source(store,tmp_path):
    from openpyxl import Workbook,load_workbook
    from excel.file_io import process_file,HEADERS
    from excel.simple_review import build_review_queue,refresh_review_row
    from excel.output_update import apply_decisions
    path=tmp_path/'Исходный.xlsx';wb=Workbook();ws=wb.active
    ws.append(['Код Автодокс','Наименование','Завод'])
    ws.append(['1','Панель ПКМ-ТСТ-1',''])
    ws.append(['2','Панель ПКМ-ТСТ-2',''])
    wb.save(path);wb.close();before=path.read_bytes()
    base=RemovesFamilies()
    out,report=process_file(path,tmp_path,OperatorExpertAnonymizer(store.snapshot(),base,True),knowledge_store=store)
    rows=build_review_queue(store,[report['session']],base)
    assert len(rows)==2 and 'ПКМ-ТСТ' not in rows[1]['automatic']
    store.decide(rows[0],rows[0]['source'],'Оставить как в исходном')
    current=refresh_review_row(store,rows[1],base)
    assert current['automatic']==current['source'] and current['provenance']['series_kept']==['ПКМ-ТСТ']
    store.decide(current,current['automatic'],'Правильно')
    apply_decisions(out,[dict(row=rows[0],final=rows[0]['source'],action='Оставить как в исходном'),dict(row=current,final=current['automatic'],action='Правильно')])
    wb=load_workbook(out);ws=wb.active;c=[x.value for x in ws[1]].index(HEADERS[1])+1
    assert ws.cell(2,c).value==rows[0]['source'] and ws.cell(3,c).value==rows[1]['source'];wb.close()
    assert path.read_bytes()==before


def test_export_import_can_disable_learned_series(store,tmp_path):
    from openpyxl import load_workbook
    from excel.knowledge_edit_io import export_editable,commit_editable
    record(store,'Модуль МОС-1')
    path=tmp_path/'База.xlsx';export_editable(path,store)
    wb=load_workbook(path);ws=wb['Управление'];cols={c.value:c.column for c in ws[1]}
    n=next(i for i in range(2,ws.max_row+1) if str(ws.cell(i,cols['Ключ']).value).startswith('series:'))
    ws.cell(n,cols['Действие'],'ОТКЛЮЧИТЬ');ws.cell(n,cols['Причина изменения'],'Не обобщать');wb.save(path);wb.close()
    assert commit_editable(path,store)['applied']==1
    assert not result(store,'Модуль МОС-2')['knowledge'].get('series_kept')
    assert result(store,'Модуль МОС-1',code='1')['text']=='Модуль МОС-1'


def test_changed_source_same_code_still_requires_review(store):
    record(store,'Модуль МОС-1')
    got=result(store,'Модуль МОС-2',code='1')
    assert got['text']=='Модуль МОС-2'
    assert got['status']=='ЖЁЛТЫЙ' and not got['knowledge'].get('exact_applied')


def test_color_spans_follow_edit_and_restore():
    from excel.review_display import change_spans,restored_spans
    source='Панель ПКМ-ТСТ-3 IP66';automatic='Панель IP66'
    deleted,added=change_spans(source,automatic)
    assert ''.join(source[a:b] for a,b in deleted).strip()=='ПКМ-ТСТ-3' and not added
    assert not change_spans(source,source)[0]
    assert ''.join(source[a:b] for a,b in restored_spans(source,automatic,source)).strip()=='ПКМ-ТСТ-3'
    final=source+' в комплекте'
    assert ''.join(final[a:b] for a,b in change_spans(source,final)[1]).strip()=='в комплекте'


def test_sync_notices_distinguish_offline_and_bad_record():
    from excel.sync_feedback import feedback
    offline=feedback('Папка',dict(online=False,pending=1,errors=['Отказано в доступе']))
    assert 'Нет доступа' in offline['title'] and 'Отказано в доступе' in offline['details']
    repeated=feedback('Папка',dict(online=False,pending=2,errors=['Отказано в доступе: другой файл']))
    assert offline['signature']==repeated['signature']
    damaged=feedback('Папка',dict(online=True,pending=0,errors=['Повреждён журнал']))
    assert 'записи с ошибками' in damaged['title'] and damaged['signature']!=offline['signature']
    assert feedback('Папка',dict(online=True,pending=0,errors=[]))['ok']
