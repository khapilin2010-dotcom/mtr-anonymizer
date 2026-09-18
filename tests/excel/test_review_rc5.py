import uuid
from unittest.mock import patch
import pytest
from openpyxl import Workbook, load_workbook
from excel.simple_store import SimpleKnowledgeStore
from excel.operator_engine import OperatorExpertAnonymizer
from excel.excel_engine import Anonymizer
from excel.simple_review import build_review_queue, advance_review_queue
from excel.review_display import restore_deleted
from excel.file_io import process_file, HEADERS
from excel.semantics import features

@pytest.fixture
def store(tmp_path):
    return SimpleKnowledgeStore(tmp_path/'local',tmp_path/'base','one-engineer')

@pytest.fixture(scope='module')
def base():
    return Anonymizer()


def engine(store,base):
    return OperatorExpertAnonymizer(store.snapshot(),base,True)


def confirm(store,base,n,source=None):
    source=source or f'Панель ПКМ-ТСТ-{n}'
    result=engine(store,base).anonymize(source,str(n),'')
    row=dict(id=str(uuid.uuid4()),session='manual',source=source,code=str(n),factory='',
             automatic=result['text'],classification=result['classification'],provenance=result['knowledge'])
    store.decide(row,source,'Оставить как в исходном')
    return row


@pytest.mark.parametrize('tail',['',' ТУ 1234-567-2020',' по типу ТУ 1234-567-2020',' письмо № 123 от 01.02.2020',' ИНН 1234567890'])
def test_three_positions_same_engineer_skip_only_safe_changes(store,base,tail):
    for n in (1,2):confirm(store,base,n)
    assert not engine(store,base).anonymize('Панель ПКМ-ТСТ-4'+tail,'4')['knowledge'].get('auto_reviewed')
    confirm(store,base,3)
    got=engine(store,base).anonymize('Панель ПКМ-ТСТ-4'+tail,'4')
    assert got['knowledge']['auto_reviewed'],got
    assert got['text']=='Панель ПКМ-ТСТ-4'
    assert got['knowledge']['series_confirmations']=={'ПКМ-ТСТ':3}


def test_repeat_one_case_cannot_train_itself(store,base):
    row=confirm(store,base,1)
    for _ in range(5):store.decide(row,row['source'],'Оставить как в исходном')
    got=engine(store,base).anonymize('Панель ПКМ-ТСТ-4','4')
    assert got['knowledge']['series_confirmations']=={'ПКМ-ТСТ':1}
    assert not got['knowledge'].get('auto_reviewed')


@pytest.mark.parametrize('source,factory,code',[
 ('Панель ПКМ-ТСТ-4 HAWLE-TEST9 ТУ 1234-567-2020','','4'),
 ('Панель ПКМ-ТСТ-4 ООО «Новый завод»','','4'),
 ('Панель ПКМ-ТСТ-4 XYZ-400','','4'),
 ('Насос ПКМ-ТСТ-4','','4'),
 ('Панель ПКМ-ТСТ-4','Другой завод','4'),
 ('Панель ПКМ-ТСТ-4','','1'),
])
def test_other_changes_and_changed_context_need_review(store,base,source,factory,code):
    for n in (1,2,3):confirm(store,base,n)
    got=engine(store,base).anonymize(source,code,factory)
    assert not got['knowledge'].get('auto_reviewed'),got


def test_conflict_and_disabled_rule_block_skip(store,base):
    for n in (1,2,3):confirm(store,base,n)
    key=next(k for k in store.snapshot().entries if k.startswith('series:'))
    store.control(key,True,'Не обобщать')
    assert not engine(store,base).anonymize('Панель ПКМ-ТСТ-4','4')['knowledge'].get('auto_reviewed')
    store.control(key,False,'Подтверждено')
    row=dict(id='conflict',session='manual',source='Панель ПКМ-ТСТ-5',automatic='Панель',code='5',factory='',classification=features('Панель ПКМ-ТСТ-5'))
    store.decide(row,'Панель','Правильно')
    got=engine(store,base).anonymize('Панель ПКМ-ТСТ-4','4')
    assert got['status']=='КРАСНЫЙ' and not got['knowledge'].get('auto_reviewed')


def workbook(store,base,tmp_path):
    src=tmp_path/'Исходный.xlsx';wb=Workbook();ws=wb.active
    ws.append(['Код Автодокс','Наименование','Завод'])
    for n in range(1,8):ws.append([str(n),f'Панель ПКМ-ТСТ-{n} по типу ТУ 1234-567-2020',''])
    wb.save(src);wb.close()
    out,report=process_file(src,tmp_path,engine(store,base),knowledge_store=store)
    return src,out,report


def test_third_acceptance_hides_remaining_rows_after_output_write(store,base,tmp_path):
    src,out,report=workbook(store,base,tmp_path);before=src.read_bytes()
    queue=build_review_queue(store,[report['session']],base)
    assert len(queue)==7
    for n in range(3):
        assert advance_review_queue(store,queue,n,base)==n
        store.decide(queue[n],f'Панель ПКМ-ТСТ-{n+1}','Исправить')
    event_count=len(store.events())
    assert advance_review_queue(store,queue,3,base)==7
    assert all(row.get('final')==f'Панель ПКМ-ТСТ-{i+1}' for i,row in enumerate(queue) if i>=3)
    assert len(store.events())==event_count  # No self-confirmations from automatic rows.
    wb=load_workbook(out);ws=wb.active;c=[x.value for x in ws[1]].index(HEADERS[1])+1
    assert [ws.cell(i+2,c).value for i in range(3,7)]==[f'Панель ПКМ-ТСТ-{i}' for i in range(4,8)]
    status_col=[x.value for x in ws[1]].index(HEADERS[2])+1
    assert ws.cell(5,status_col).value=='ПРОВЕРЕНО АВТОМАТИЧЕСКИ ПО БАЗЕ'
    wb.close();assert src.read_bytes()==before
    assert build_review_queue(store,[report['session']],base)==[]


def test_write_failure_must_not_hide_rows(store,base,tmp_path):
    _,_,report=workbook(store,base,tmp_path)
    queue=build_review_queue(store,[report['session']],base)
    for n in range(3):store.decide(queue[n],f'Панель ПКМ-ТСТ-{n+1}','Исправить')
    with patch('excel.simple_review.apply_decisions',side_effect=PermissionError('Excel открыт')):
        with pytest.raises(PermissionError):advance_review_queue(store,queue,3,base)
    assert all('final' not in row for row in queue[3:])
    assert advance_review_queue(store,queue,3,base)==7


def test_restore_selection_of_red_span_only_and_preserve_manual_addition():
    source='Панель ПКМ-ТСТ-3 ТУ 1234 IP66'
    final='Панель IP66 в комплекте'
    a=source.index('ПКМ');b=a+len('ПКМ-ТСТ-3')
    got=restore_deleted(source,final,(a,b))
    assert got=='Панель ПКМ-ТСТ-3 IP66 в комплекте'
    assert restore_deleted(source,got,(a,b))==got
    assert restore_deleted(source,got,(0,6))==got
    assert restore_deleted(source,got)==source+' в комплекте'


def test_restore_multiple_removed_ranges_and_partial_word():
    source='Модуль МОС-1 IP66 ТУ 1234 PN16'
    final='Модуль IP66 PN16'
    assert restore_deleted(source,final)==source
    assert restore_deleted('АБВГ','АГ',(1,2))=='АБГ'


def test_resume_rc4_rows_recomputes_even_if_database_version_unchanged(store,base):
    from excel.simple_review import refresh_review_row
    for n in (1,2,3):confirm(store,base,n)
    source='Панель ПКМ-ТСТ-4 ТУ 1234-567-2020'
    old=dict(source=source,automatic='Панель ПКМ-ТСТ-4',code='4',factory='',status='ЖЁЛТЫЙ',
             provenance=dict(version=store.snapshot().version,series_rules=['old']))
    updated=refresh_review_row(store,old,base)
    assert updated['provenance']['auto_reviewed']


def test_three_versions_of_same_code_count_as_one_position(store,base):
    for n in (1,2,3):confirm(store,base,1,source=f'Панель ПКМ-ТСТ-{n}')
    got=engine(store,base).anonymize('Панель ПКМ-ТСТ-4','4')
    assert got['knowledge']['series_confirmations']=={'ПКМ-ТСТ':1}
    assert not got['knowledge'].get('auto_reviewed')
