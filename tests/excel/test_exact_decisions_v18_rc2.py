"""End-to-end regressions for the engineer's exact decision and editable base."""
import uuid
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel.excel_engine import Anonymizer
from excel.file_io import HEADERS, process_file
from excel.knowledge import case_key, legacy_case_key, canonical, digest
from excel.knowledge_edit_io import export_editable, commit_editable, preview_editable
from excel.operator_engine import OperatorExpertAnonymizer
from excel.output_update import apply_decisions
from excel.simple_review import build_review_queue
from excel.simple_store import SimpleKnowledgeStore

MOS = 'Модуль автоматизированной технологической обвязки скважин МОС-3/1 Чертеж МОС-29.М1'


@pytest.fixture
def store(tmp_path):
    return SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'shared', 'engineer-one')


def row(source=MOS, code='631-336536', factory=''):
    return dict(id=str(uuid.uuid4()), session=str(uuid.uuid4()), source=source, code=code,
                factory=factory, automatic=source.replace('МОС', ''), status='ЖЁЛТЫЙ')


class NeverAutomatic(Anonymizer):
    def anonymize(self, *args, **kwargs):
        raise AssertionError('Automatic algorithm overwrote engineer decision')


def result(store, r, base=None):
    return OperatorExpertAnonymizer(store.snapshot(), base=base).anonymize(r['source'], r['code'], r['factory'])


def edit(path, key, action, final=None):
    wb = load_workbook(path); ws = wb['Управление']
    cols = {c.value: c.column for c in ws[1]}
    n = next(n for n in range(2, ws.max_row + 1) if ws.cell(n, cols['Ключ']).value == key)
    ws.cell(n, cols['Действие'], action)
    ws.cell(n, cols['Причина изменения'], 'Исправлено инженером')
    if final is not None:
        ws.cell(n, cols['Новое решение'], final)
    wb.save(path); wb.close()


@pytest.mark.parametrize('action,final', [
    ('Оставить как в исходном', MOS),
    ('Исправить', MOS + ' в комплекте'),
    ('Правильно', MOS.replace('МОС', '')),
])
def test_three_decisions_survive_repeat_and_new_computer(store, tmp_path, action, final):
    r = row(); store.decide(r, final, action); store.sync(auto_backup=False)
    fresh = SimpleKnowledgeStore(tmp_path / 'another-machine', store.shared, 'engineer-two')
    fresh.sync(auto_backup=False)
    for _ in range(2):
        got = result(fresh, r, NeverAutomatic())
        assert got['text'] == final
        assert got['knowledge']['exact_applied'] is True


def test_two_sources_same_code_retain_both_decisions(store):
    a, b = row(), row(MOS + ' комплект 2')
    store.decide(a, a['source'], 'Оставить как в исходном')
    before = result(store, b)
    assert before['status'] == 'ЖЁЛТЫЙ'
    assert before['knowledge']['exact_applied'] is False
    store.decide(b, b['source'] + ' проверено')
    assert result(store, a, NeverAutomatic())['text'] == a['source']
    assert result(store, b, NeverAutomatic())['text'] == b['source'] + ' проверено'
    assert case_key(a['code'], a['source']) != case_key(b['code'], b['source'])


def test_changed_factory_requires_review(store):
    a, b = row(factory='Первый завод'), row(factory='Другой завод')
    store.decide(a, a['source'])
    got = result(store, b)
    assert got['status'] == 'ЖЁЛТЫЙ'
    assert not got['knowledge'].get('exact_applied')


def test_context_change_not_silenced_by_general_auto_mode(store):
    a = row(); store.decide(a, a['source'])
    for mode in (False, True):
        got = OperatorExpertAnonymizer(store.snapshot(), auto_apply_confirmed=mode).anonymize(MOS + ' другой', a['code'])
        assert got['status'] == 'ЖЁЛТЫЙ'
        assert not got['knowledge'].get('exact_applied')


def test_export_has_russian_status_and_searchable_metadata(store, tmp_path):
    r = row(); store.decide(r, MOS, 'Оставить как в исходном')
    path = tmp_path / 'knowledge.xlsx'; export_editable(path, store)
    wb = load_workbook(path); ws = wb['Управление']; headers = [c.value for c in ws[1]]
    assert set(('Код Автодокс', 'Исходное наименование', 'Завод', 'Текущее решение', 'Тип решения',
                'Дата решения', 'Инженер', 'Статус', 'Независимых инженеров')) <= set(headers)
    assert ws.cell(2, headers.index('Статус') + 1).value == 'ДЕЙСТВУЕТ'
    assert ws.cell(2, headers.index('Код Автодокс') + 1).value == r['code']
    assert ws.cell(2, headers.index('Исходное наименование') + 1).value == MOS
    assert ws.auto_filter.ref
    wb.close()


def test_excel_correction_by_another_engineer_is_authoritative(store, tmp_path):
    r = row(); store.decide(r, MOS); store.sync(auto_backup=False)
    other = SimpleKnowledgeStore(tmp_path / 'other', store.shared, 'engineer-two')
    path = tmp_path / 'edit.xlsx'; export_editable(path, other)
    key = case_key(r['code'], MOS)
    edit(path, key, 'ИЗМЕНИТЬ', MOS + ' исправлено')
    assert commit_editable(path, other)['applied'] == 1
    store.sync(auto_backup=False)
    assert result(store, r, NeverAutomatic())['text'] == MOS + ' исправлено'
    assert len(store.snapshot().entries[key]['history']) >= 2
    assert commit_editable(path, other)['applied'] == 0


def test_disable_enable_and_reconfirm(store, tmp_path):
    r = row(); store.decide(r, MOS)
    key = case_key(r['code'], MOS); path = tmp_path / 'edit.xlsx'
    for action in ('ОТКЛЮЧИТЬ', 'ВКЛЮЧИТЬ', 'ОТКЛЮЧИТЬ'):
        export_editable(path, store); edit(path, key, action)
        assert commit_editable(path, store)['applied'] == 1
        assert bool(result(store, r).get('knowledge', {}).get('exact_applied')) == (action == 'ВКЛЮЧИТЬ')
    store.decide(r, MOS + ' проверено')
    assert result(store, r, NeverAutomatic())['text'] == MOS + ' проверено'


def test_stale_status_change_is_rejected(store, tmp_path):
    r = row(); store.decide(r, MOS); key = case_key(r['code'], MOS)
    path = tmp_path / 'edit.xlsx'; export_editable(path, store)
    edit(path, key, 'ИЗМЕНИТЬ', MOS + ' старое исправление')
    store.control(key, True, 'Отключено после выгрузки')
    report = commit_editable(path, store)
    assert report['applied'] == 0 and report['issues']
    assert store.snapshot().entries[key]['status'] == 'DISABLED'


def test_protected_parameter_cannot_be_lost_via_excel(store, tmp_path):
    r = row(MOS + ' IP66 УХЛ1'); store.decide(r, r['source'])
    key = case_key(r['code'], r['source']); path = tmp_path / 'edit.xlsx'
    export_editable(path, store); edit(path, key, 'ИЗМЕНИТЬ', MOS)
    report = commit_editable(path, store)
    assert report['applied'] == 0 and report['issues']
    assert result(store, r)['text'] == r['source']


def test_legacy_journal_projects_context_without_rewriting(store):
    a, b = row(), row(MOS + ' другой вариант')
    first = store.decide(a, MOS)
    second = store.decide(b, b['source'])
    first['key'] = second['key'] = legacy_case_key(a['code'], MOS)
    second['supersedes'] = [first['id']]
    with store.db() as db:
        for event in (first, second):
            db.execute('UPDATE events SET body=? WHERE id=?', (canonical(event), event['id']))
    store.sync(auto_backup=False)
    before = {p: p.read_bytes() for p in (store.shared / 'events').glob('*/*.json')}
    for r in (a, b):
        assert result(store, r, NeverAutomatic())['text'] == r['source']
    fresh = SimpleKnowledgeStore(store.local / 'fresh', store.shared, 'engineer-two')
    assert not fresh.sync(auto_backup=False)['errors']
    assert result(fresh, a, NeverAutomatic())['text'] == MOS
    store.decide(a, MOS + ' новая правка'); store.sync(auto_backup=False)
    assert result(store, a, NeverAutomatic())['text'] == MOS + ' новая правка'
    assert all(p.read_bytes() == content for p, content in before.items())


def create_session(store, tmp_path):
    source = tmp_path / 'исходник.xlsx'; wb = Workbook(); ws = wb.active
    ws.append(['Код Автодокс', 'Наименование', 'Завод'])
    ws.append(['631-336536', MOS, ''])
    ws.append(['OTHER', MOS + ' второй', ''])
    wb.save(source); wb.close()
    before = source.read_bytes()
    out, report = process_file(source, tmp_path, OperatorExpertAnonymizer(store.snapshot()), knowledge_store=store)
    return source, before, out, report


def final_cell(path, n=2):
    wb = load_workbook(path); ws = wb.active
    col = [c.value for c in ws[1]].index(HEADERS[1]) + 1
    value = ws.cell(n, col).value; wb.close(); return value


def test_crash_recovery_writes_decision_before_skipping_and_preserves_source(store, tmp_path):
    source, before, output, report = create_session(store, tmp_path)
    r = store.session_rows(report['session'])[0]
    store.decide(r, MOS, 'Оставить как в исходном')  # simulate crash before output flush
    assert build_review_queue(store, [report['session']]) is not None
    assert final_cell(output) == MOS
    assert all(item['id'] != r['id'] for item in build_review_queue(store, [report['session']]))
    again, _ = process_file(source, tmp_path, OperatorExpertAnonymizer(store.snapshot()), knowledge_store=store)
    assert final_cell(again) == MOS
    assert source.read_bytes() == before
    assert again != output and output != source


def test_output_failure_cannot_hide_saved_decision(store, tmp_path):
    _, _, output, report = create_session(store, tmp_path)
    r = store.session_rows(report['session'])[0]; store.decide(r, MOS)
    Path(output).unlink()
    with pytest.raises(FileNotFoundError):
        build_review_queue(store, [report['session']])


def test_original_and_hardlink_cannot_be_written(store, tmp_path):
    source, before, output, report = create_session(store, tmp_path)
    r = store.session_rows(report['session'])[0]
    link = tmp_path / 'hardlink.xlsx'; link.hardlink_to(source)
    for target in (source, link):
        with pytest.raises(ValueError, match='Исходный файл'):
            apply_decisions(target, [dict(row=r, final=MOS, action='Правильно')])
    assert source.read_bytes() == before


def test_sorted_output_updates_by_row_identity(store, tmp_path):
    source, before, output, report = create_session(store, tmp_path)
    r = store.session_rows(report['session'])[0]
    wb = load_workbook(output); ws = wb.active
    a, b = [[c.value for c in ws[n]] for n in (2, 3)]
    for n, values in ((2, b), (3, a)):
        for col, value in enumerate(values, 1):
            ws.cell(n, col).value = value
    wb.save(output); wb.close()
    apply_decisions(output, [dict(row=r, final=MOS + ' исправлено', action='Исправить')])
    assert final_cell(output, 3) == MOS + ' исправлено'
    assert source.read_bytes() == before


def test_disabling_case_also_retires_derived_rules(store):
    r = row('Клапан XYZ-500 DN50', factory='Тестовый завод')
    store.decide(r, 'Клапан DN50')
    key = case_key(r['code'], r['source'], r['factory'])
    store.control(key, True, 'Ошибочное решение')
    assert all(e['status'] == 'DISABLED' for e in store.snapshot().entries.values())


def test_same_result_counts_independent_engineers(store, tmp_path):
    r = row(); store.decide(r, MOS); store.sync(auto_backup=False)
    other = SimpleKnowledgeStore(tmp_path / 'other', store.shared, 'engineer-two')
    other.sync(auto_backup=False); other.decide(r, MOS); other.sync(auto_backup=False)
    key = case_key(r['code'], MOS)
    assert other.snapshot().entries[key]['confirmations'] == 2
    assert other.decide(r, MOS) is None
    other.decide(r, MOS + ' исправлено')
    assert result(other, r, NeverAutomatic())['text'] == MOS + ' исправлено'


def test_edit_derived_rule_modifies_same_rule(store, tmp_path):
    r = row('Клапан XYZ-500 DN50', factory='Тестовый завод')
    r['automatic'] = r['source']
    store.decide(r, 'Клапан DN50')
    key = next(k for k, e in store.snapshot().entries.items() if e['kind'] == 'rule')
    path = tmp_path / 'rule.xlsx'; export_editable(path, store)
    edit(path, key, 'ИЗМЕНИТЬ', 'ОСТАВИТЬ')
    assert commit_editable(path, store)['applied'] == 1
    assert store.snapshot().entries[key]['value'] == 'KEEP'
    for action in ('ОТКЛЮЧИТЬ', 'ВКЛЮЧИТЬ'):
        export_editable(path, store); edit(path, key, action)
        assert commit_editable(path, store)['applied'] == 1
    assert store.snapshot().entries[key]['value'] == 'KEEP'
    assert store.snapshot().entries[key]['status'] == 'ACTIVE'


def test_legacy_disable_can_be_enabled_after_migration(store):
    r = row(); event = store.decide(r, MOS); old_key = legacy_case_key(r['code'], MOS)
    event['key'] = old_key
    with store.db() as db:
        db.execute('UPDATE events SET body=? WHERE id=?', (canonical(event), event['id']))
    store.control(old_key, True, 'Старое отключение')
    key = case_key(r['code'], MOS)
    assert store.snapshot().entries[key]['status'] == 'DISABLED'
    store.control(key, False, 'Исправлено в новой версии')
    assert result(store, r, NeverAutomatic())['text'] == MOS
