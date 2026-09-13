import uuid

from openpyxl import Workbook, load_workbook

from excel.knowledge import Snapshot, case_key
from excel.mapped_import import header_candidates, guess_mapping, preview_mapped, commit_mapped
from excel.quality_tools import (
    export_manager_report, is_unknown_row, knowledge_audit, previous_decision,
    removed_ratio, row_risks, same_case,
)
from excel.simple_store import SimpleKnowledgeStore


def _event(source, final, user='ENGINEER\\one', code='', factory='', timestamp='2026-09-13T12:00:00+00:00'):
    key = case_key(code, source, factory)
    return dict(schema=1, store_id=str(uuid.uuid4()), id=str(uuid.uuid4()), user=user,
                timestamp=timestamp, version='test', kind='case', key=key, value=final,
                source=source, code=code, factory=factory, automatic=source,
                row_id=str(uuid.uuid4()), session=str(uuid.uuid4()), action='Исправить',
                supersedes=[])


def test_removed_ratio_and_dangerous_short_result():
    assert removed_ratio('1234567890', '12') >= 0.7
    row = {'source': 'Клапан DN50 IP66 TESTBRAND', 'automatic': 'Клапан', 'code': '', 'factory': '', 'status': 'ЖЁЛТЫЙ'}
    risks = row_risks(row)
    assert any('защищ' in x.lower() for x in risks)
    assert any('%' in x for x in risks)


def test_previous_decision_comparison_and_training_unknown():
    source = 'Насос TESTBRAND DN50'
    snapshot = Snapshot([_event(source, 'Насос DN50')])
    row = {'source': source, 'automatic': source, 'code': '', 'factory': '', 'status': 'ЖЁЛТЫЙ'}
    previous = previous_decision(row, snapshot)
    assert previous['value'] == 'Насос DN50'
    assert not is_unknown_row(row, snapshot)
    assert any('ранее подтверждённого' in x.lower() for x in row_risks(row, snapshot))
    other = dict(row, source='Новая неизвестная позиция')
    assert is_unknown_row(other, snapshot)


def test_same_case_requires_source_factory_and_automatic_match():
    a = {'source': 'Насос X', 'factory': 'Завод', 'automatic': 'Насос'}
    b = {'source': ' Насос X ', 'factory': 'Завод', 'automatic': 'Насос'}
    c = dict(b, factory='Другой завод')
    assert same_case(a, b)
    assert not same_case(a, c)


def test_knowledge_audit_finds_disputed_case():
    source = 'Насос TESTBRAND DN50'
    events = [_event(source, 'Насос DN50', 'ENGINEER\\one'),
              _event(source, source, 'ENGINEER\\two', timestamp='2026-09-13T12:01:00+00:00')]
    snapshot = Snapshot(events)
    issues = knowledge_audit(snapshot)
    assert any(x['severity'] == 'КРАСНЫЙ' and 'Противореч' in x['issue'] for x in issues)


def test_arbitrary_headers_can_be_mapped_and_imported(tmp_path):
    shared = tmp_path / 'MTR_Knowledge'
    store = SimpleKnowledgeStore(tmp_path / 'local', shared, user='ENGINEER\\one')
    path = tmp_path / 'legacy_weird.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Старая база'
    ws.append(['Внутренний номер', 'Текст до', 'Кто сделал', 'Текст после'])
    ws.append(['A-1', 'Насос TESTBRAND DN50', 'Тестовый завод', 'Насос DN50'])
    wb.save(path); wb.close()

    candidates = header_candidates(path)
    candidate = next(x for x in candidates if x['row'] == 1)
    guessed = guess_mapping(candidate['headers'])
    assert guessed['source'] is None  # deliberately non-standard headings
    mapping = {'code': 0, 'source': 1, 'factory': 2, 'final': 3}
    plan = preview_mapped(path, store, candidate, mapping)
    assert plan['ready'] == 1
    report = commit_mapped(path, store, candidate, mapping)
    assert report['accepted'] == 1
    assert store.snapshot().entries[case_key('A-1', 'Насос TESTBRAND DN50', 'Тестовый завод')]['value'] == 'Насос DN50'


def test_manager_report_contains_quality_sheets(tmp_path):
    store = SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'MTR_Knowledge', user='ENGINEER\\one')
    row = dict(id=str(uuid.uuid4()), session='manual', code='', source='Насос TESTBRAND DN50',
               factory='', automatic='Насос TESTBRAND DN50', classification=None,
               explicit_feedback=True, provenance={})
    store.decide(row, 'Насос DN50', 'Исправить'); store.sync()
    path = tmp_path / 'report.xlsx'
    export_manager_report(path, store)
    wb = load_workbook(path, read_only=True)
    assert {'Сводка', 'Сеансы', 'Контроль базы', 'Производители'} <= set(wb.sheetnames)
    wb.close()
