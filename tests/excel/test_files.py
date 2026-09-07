import csv
import hashlib
import io
from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel.excel_engine import Anonymizer
from excel.file_io import HEADERS, process_file


@pytest.fixture(scope='module')
def az():
    return Anonymizer()


def test_source_cells_formulas_layout_and_report(tmp_path, az):
    folder = tmp_path / 'Русская папка с пробелами'
    folder.mkdir()
    src = folder / 'данные.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Готово'
    ws.merge_cells('A1:C1'); ws['A1'] = 'Синтетические данные'
    ws.append(['Код ресурса', 'Наименование', 'Производитель'])
    ws.append(['001234', 'Клапан Унипол IP66', 'Исходный поставщик'])
    ws.append(['000002', '=B3', ''])
    ws['A3'].font = Font(bold=True, color='112233')
    ws['A3'].number_format = '@'
    ws.row_dimensions[3].height = 44
    extra = wb.create_sheet('Справка'); extra['A1'] = '=1+2'
    wb.save(src); before = src.read_bytes()
    dst, report = process_file(src, folder, az)
    result = load_workbook(dst)
    # Compare to the serialized source, including Excel's empty-cell representation.
    ws = load_workbook(src)['Готово']
    assert src.read_bytes() == before
    for row in range(1, 5):
        for col in range(1, 4):
            assert result['Готово'].cell(row, col).value == ws.cell(row, col).value
            assert result['Готово'].cell(row, col).style_id == ws.cell(row, col).style_id
    assert result['Готово']['F4'].value.startswith('ЖЁЛТЫЙ: формула')
    assert result['Справка']['A1'].value == '=1+2'
    assert list(result['Готово'].merged_cells.ranges) == list(ws.merged_cells.ranges)
    assert result['Готово'].row_dimensions[3].height == 44
    assert [c.value for c in result['Готово'][2]][3:] == list(HEADERS)
    assert result['Готово']['E3'].value == 'Клапан IP66'
    assert report['skipped_formulas'] == 1
    assert report['rows'] == 1


@pytest.mark.parametrize('encoding,delimiter', [('utf-8-sig',';'), ('cp1251',';'), ('utf-8',','), ('utf-8','\t')])
def test_csv_multiline_and_original_columns(tmp_path, az, encoding, delimiter):
    src = tmp_path / 'исходник.csv'
    source = [['Код Автодокс','Наименование','Производитель'],
              ['001234','Клапан Унипол\nIP66','Прежний поставщик'],
              ['000002','Кабель ТУ 123-45 220 В','']]
    with src.open('w', encoding=encoding, newline='') as f:
        csv.writer(f, delimiter=delimiter).writerows(source)
    before = src.read_bytes()
    dst, report = process_file(src, tmp_path, az)
    with dst.open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f, delimiter=delimiter))
    assert [row[:3] for row in rows] == source
    assert rows[1][4] == 'Клапан IP66'
    assert rows[2][4] == 'Кабель 220 В'
    assert report['rows'] == 2
    assert src.read_bytes() == before


def test_existing_outputs_never_overwritten(tmp_path, az):
    src = tmp_path / 'in.csv'; src.write_text('Код;Наименование\n1;Клапан Унипол IP66',encoding='utf-8')
    first, _ = process_file(src, tmp_path, az)
    before = first.read_bytes()
    second, _ = process_file(src, tmp_path, az)
    assert first != second
    assert first.read_bytes() == before


def test_bad_file_cleans_partial_output(tmp_path, az):
    src = tmp_path / 'broken.xlsx'; src.write_bytes(b'not a workbook')
    with pytest.raises(zipfile.BadZipFile):
        process_file(src, tmp_path, az)
    assert list(tmp_path.iterdir()) == [src]


def test_xls_preserves_input_and_adds_four_columns(tmp_path, az):
    import xlwt, xlrd
    src = tmp_path / 'legacy.xls'
    wb = xlwt.Workbook(); ws = wb.add_sheet('Готово')
    for i,v in enumerate(['Код','Наименование','Производитель']): ws.write(0,i,v)
    for i,v in enumerate(['000001','Клапан Унипол IP66','Исходный поставщик']): ws.write(1,i,v)
    wb.save(str(src)); original = src.read_bytes()
    dst, report = process_file(src,tmp_path,az)
    out = xlrd.open_workbook(str(dst)).sheet_by_index(0)
    assert out.row_values(1)[:3] == ['000001','Клапан Унипол IP66','Исходный поставщик']
    assert out.row_values(0)[3:] == list(HEADERS)
    assert out.cell_value(1,4) == 'Клапан IP66'
    assert report['rows'] == 1
    assert src.read_bytes() == original


def test_xls_formula_cannot_be_silently_destroyed(tmp_path, az):
    import xlwt
    src = tmp_path / 'formula.xls'
    wb = xlwt.Workbook(); ws = wb.add_sheet('Готово')
    ws.write(0,0,'Код');ws.write(0,1,'Наименование')
    ws.write(1,0,'1');ws.write(1,1,'Клапан Унипол')
    ws.write(1,2,xlwt.Formula('1+2'));wb.save(str(src))
    original = src.read_bytes()
    with pytest.raises(ValueError,match='XLS содержит формулы'):
        process_file(src,tmp_path,az)
    assert src.read_bytes() == original


def test_xlsm_vba_archive_preserved(tmp_path, az):
    src = tmp_path / 'macro.xlsm'
    wb = Workbook();ws=wb.active;ws.append(['Код','Наименование']);ws.append(['1','Клапан Унипол IP66']);wb.save(src)
    # Synthetic opaque macro payload: preservation check, never executed.
    payload = b'SYNTHETIC VBA PAYLOAD - DO NOT EXECUTE'
    with zipfile.ZipFile(src,'a') as z:
        z.writestr('xl/vbaProject.bin',payload)
    dst,report=process_file(src,tmp_path,az)
    assert dst.suffix == '.xlsm'
    with zipfile.ZipFile(dst) as z:
        assert z.read('xl/vbaProject.bin') == payload
    assert report['rows'] == 1


def test_only_preferred_sheet_processed(tmp_path, az):
    src=tmp_path/'sheets.xlsx';wb=Workbook();wb.active.title='выборка_оборудования_все'
    for title in ['Выборка оборудования','Готово']: wb.create_sheet(title)
    for ws in wb:
        ws.append(['Код','Наименование']);ws.append(['1','Клапан Унипол IP66'])
    wb.save(src);dst,report=process_file(src,tmp_path,az);out=load_workbook(dst)
    assert report['sheets']==1
    assert out['Готово'].max_column==6
    assert out['Выборка оборудования'].max_column==2
    assert out['выборка_оборудования_все'].max_column==2


def test_generated_equals_text_is_not_formula(tmp_path,az):
    src=tmp_path/'literal.xlsx';wb=Workbook();ws=wb.active
    ws.append(['Код','Наименование']);ws.append(['1','=Клапан Унипол IP66']);ws['B2'].data_type='s';wb.save(src)
    dst,_=process_file(src,tmp_path,az);cell=load_workbook(dst).active['D2']
    assert cell.data_type=='s'
    assert cell.value=='=Клапан IP66'
