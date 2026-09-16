import uuid

import pytest
from openpyxl import Workbook, load_workbook

from excel.excel_engine import Anonymizer
from excel.file_io import HEADERS
from excel.knowledge import case_key
from excel.knowledge_edit_io import export_editable, commit_editable
from excel.review_io import REVIEW_HEADERS
from excel.output_update import apply_decisions
from excel.operator_engine import OperatorExpertAnonymizer
from excel.simple_store import SimpleKnowledgeStore


class StaticMustNotRun(Anonymizer):
    def anonymize(self, *args, **kwargs):
        raise AssertionError('static anonymizer must not run for an exact engineer decision')


def _row(source, automatic, code='631-336536', factory=''):
    return dict(
        id=str(uuid.uuid4()), session=str(uuid.uuid4()), row=12, sheet='Лист1',
        code=code, source=source, factory=factory, automatic=automatic,
        status='ЖЁЛТЫЙ', classification=None, provenance={}, removed=['МОС']
    )


def test_engineer_keep_original_wins_on_next_run(tmp_path):
    store = SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'MTR_Knowledge', user='ENGINEER\\one')
    source = 'Модуль автоматизированной технологической обвязки скважин МОС-3/1 Чертеж МОС-29.М1'
    automatic = 'Модуль автоматизированной технологической обвязки скважин 3/1 Чертеж -29.М1'
    row = _row(source, automatic)
    store.decide(row, source, 'Оставить как в исходном')
    store.sync(auto_backup=False)

    result = OperatorExpertAnonymizer(store.snapshot(), auto_apply_confirmed=False).anonymize(
        source, row['code'], row['factory'])
    assert result['text'] == source
    assert result['status'] == 'ЗЕЛЁНЫЙ'
    assert result['knowledge']['status'] in ('ACTIVE', 'TRUSTED')
    assert result['knowledge']['exact_applied'] is True


def test_exact_decision_bypasses_static_anonymizer_entirely(tmp_path):
    store = SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'MTR_Knowledge', user='ENGINEER\\one')
    source = 'Модуль МОС-3/1 Чертеж МОС-29.М1 Масса-16,2 т.'
    automatic = 'Модуль 3/1 Чертеж -29.М1 Масса-16,2 т.'
    row = _row(source, automatic)
    store.decide(row, source, 'Оставить как в исходном')
    store.sync(auto_backup=False)

    result = OperatorExpertAnonymizer(
        store.snapshot(), base=StaticMustNotRun(), auto_apply_confirmed=False
    ).anonymize(source, row['code'], row['factory'])
    assert result['text'] == source
    assert result['knowledge']['source'] == 'Точное решение инженера'
    assert result['knowledge']['exact_applied'] is True


def test_same_code_but_changed_source_is_not_silently_treated_as_exact(tmp_path):
    store = SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'MTR_Knowledge', user='ENGINEER\\one')
    source = 'Модуль МОС-3/1 Чертеж МОС-29.М1'
    row = _row(source, 'Модуль 3/1 Чертеж -29.М1')
    store.decide(row, source, 'Оставить как в исходном')
    store.sync(auto_backup=False)

    changed_source = source + ' комплект 2'
    with pytest.raises(AssertionError, match='static anonymizer'):
        OperatorExpertAnonymizer(
            store.snapshot(), base=StaticMustNotRun(), auto_apply_confirmed=False
        ).anonymize(changed_source, row['code'], row['factory'])


def test_engineer_correction_wins_on_next_run(tmp_path):
    store = SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'MTR_Knowledge', user='ENGINEER\\one')
    source = 'Модуль МОС-3/1 Чертеж МОС-29.М1 Масса-16,2 т.'
    automatic = 'Модуль 3/1 Чертеж -29.М1 Масса-16,2 т.'
    corrected = 'Модуль МОС-3/1 Чертеж МОС-29.М1 Масса-16,2 т.'
    row = _row(source, automatic)
    store.decide(row, corrected, 'Исправить')
    store.sync(auto_backup=False)
    result = OperatorExpertAnonymizer(store.snapshot(), auto_apply_confirmed=False).anonymize(
        source, row['code'], row['factory'])
    assert result['text'] == corrected
    assert result['knowledge']['exact_applied'] is True


def test_apply_decision_updates_generated_xlsx(tmp_path):
    path = tmp_path / 'result.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Выборка'
    original_headers = ['Код Автодокс', 'Наименование', 'Производитель']
    ws.append(original_headers + list(HEADERS) + list(REVIEW_HEADERS))
    source = 'Модуль МОС-3/1 Чертеж МОС-29.М1 Масса-16,2 т.'
    automatic = 'Модуль 3/1 Чертеж -29.М1 Масса-16,2 т.'
    ws.append(['631-336536', source, '', '', automatic, 'ЖЁЛТЫЙ', 'МОС', '', '', '', '', 'row-id', 'session-id'])
    wb.save(path)
    wb.close()

    row = dict(row=2, sheet='Выборка', source=source, code='631-336536', factory='')
    apply_decisions(path, [{'row': row, 'final': source, 'action': 'Оставить как в исходном'}])

    wb = load_workbook(path)
    ws = wb['Выборка']
    headers = [cell.value for cell in ws[1]]
    final_col = headers.index(HEADERS[1]) + 1
    status_col = headers.index(HEADERS[2]) + 1
    action_col = headers.index(REVIEW_HEADERS[0]) + 1
    assert ws.cell(2, final_col).value == source
    assert ws.cell(2, status_col).value == 'ПРОВЕРЕНО ИНЖЕНЕРОМ: Оставить как в исходном'
    assert ws.cell(2, action_col).value == 'Оставить как в исходном'
    wb.close()


def test_exported_knowledge_can_be_edited_and_loaded_back(tmp_path):
    store = SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'MTR_Knowledge', user='ENGINEER\\one')
    source = 'Модуль МОС-3/1 Чертеж МОС-29.М1'
    row = _row(source, 'Модуль 3/1 Чертеж -29.М1')
    store.decide(row, source, 'Оставить как в исходном')
    store.sync(auto_backup=False)

    path = tmp_path / 'knowledge.xlsx'
    export_editable(path, store)
    wb = load_workbook(path)
    ws = wb['Управление']
    headers = [c.value for c in ws[1]]
    key_col = headers.index('Ключ') + 1
    new_col = headers.index('Новое решение') + 1
    action_col = headers.index('Действие') + 1
    reason_col = headers.index('Причина изменения') + 1
    target_key = case_key(row['code'], source, row['factory'])
    target_row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, key_col).value == target_key)
    changed = source + ' ИСПРАВЛЕНО'
    ws.cell(target_row, new_col, changed)
    ws.cell(target_row, action_col, 'CHANGE')
    ws.cell(target_row, reason_col, 'тест исправления')
    wb.save(path)
    wb.close()

    report = commit_editable(path, store)
    assert report['applied'] == 1
    assert store.snapshot().entries[target_key]['value'] == changed
