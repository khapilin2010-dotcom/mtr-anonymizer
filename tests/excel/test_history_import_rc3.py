"""Previously reviewed selections must become exact, traceable decisions."""
import csv
import uuid
import pytest
from openpyxl import Workbook, load_workbook

from excel.excel_engine import Anonymizer
from excel.history_import import best_candidate, preview_selection, commit_selection
from excel.mapped_import import header_candidates, guess_mapping
from excel.knowledge import case_key
from excel.simple_store import SimpleKnowledgeStore
from excel.operator_engine import OperatorExpertAnonymizer
from excel.file_io import process_file, HEADERS

MOS = 'Модуль автоматизированной технологической обвязки скважин МОС-3/1 Чертеж МОС-29.М1'
HEAD = ['Код Автодокс', 'Исходное наименование', 'Обезличенное наименование', 'Завод']


@pytest.fixture
def store(tmp_path):
    return SimpleKnowledgeStore(tmp_path / 'local', tmp_path / 'shared', 'engineer')


def workbook(tmp_path, rows=None, suffix='.xlsx'):
    path = tmp_path / ('Проверенная выборка' + suffix)
    rows = [HEAD] + (rows if rows is not None else [['00001', MOS, MOS, 'Завод А']])
    if suffix == '.csv':
        with path.open('w', encoding='utf-8-sig', newline='') as handle:
            csv.writer(handle, delimiter=';').writerows(rows)
    elif suffix == '.xls':
        import xlwt
        wb = xlwt.Workbook(); ws = wb.add_sheet('Выборка')
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                ws.write(i, j, value)
        wb.save(str(path))
    else:
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path); wb.close()
    return path


def preview(path, store):
    candidates = header_candidates(path)
    candidate = candidates[best_candidate(candidates)]
    return preview_selection(path, store, candidate, guess_mapping(candidate['headers']))


def decide(store, final=MOS, code='00001', factory='Завод А'):
    return store.decide(dict(id=str(uuid.uuid4()), session='test', source=MOS, code=code,
                            factory=factory, automatic=MOS), final)


class NeverAutomatic(Anonymizer):
    def anonymize(self, *args, **kwargs):
        raise AssertionError('Exact import must bypass automation')


@pytest.mark.parametrize('suffix', ['.xlsx', '.xlsm', '.xls', '.csv'])
def test_import_repeat_exact_and_unchanged(store, tmp_path, suffix):
    path = workbook(tmp_path, suffix=suffix); original = path.read_bytes()
    plan = preview(path, store)
    assert plan['new'] == 1 and not plan['issues'] and not store.snapshot().entries
    assert commit_selection(path, store, plan)['accepted'] == 1
    got = OperatorExpertAnonymizer(store.snapshot(), base=NeverAutomatic()).anonymize(MOS, '00001', 'Завод А')
    assert got['text'] == MOS and got['knowledge']['exact_applied']
    count = len(store.events())
    repeated = preview(path, store)
    assert repeated['repeated'] == 1 and repeated['ready'] == 0
    assert commit_selection(path, store, repeated)['accepted'] == 0
    assert len(store.events()) == count
    assert path.read_bytes() == original
    history = store.snapshot().entries[case_key('00001', MOS, 'Завод А')]['history']
    assert history[0]['provenance']['reviewed'] is True
    assert history[0]['provenance']['file_sha256'] == plan['file_hash']


def test_manual_result_writes_to_output_and_context_change_reviews(store, tmp_path):
    final = MOS + ' в комплекте'
    path = workbook(tmp_path, [['00001', MOS, final, 'Завод А']])
    commit_selection(path, store, preview(path, store))
    src = tmp_path / 'Новая выборка.xlsx'; wb = Workbook(); ws = wb.active
    ws.append(['Код Автодокс', 'Наименование', 'Завод'])
    ws.append(['00001', MOS, 'Завод А']); wb.save(src); wb.close(); before = src.read_bytes()
    out, _ = process_file(src, tmp_path, OperatorExpertAnonymizer(store.snapshot()), knowledge_store=store)
    wb = load_workbook(out); ws = wb.active
    assert ws.cell(2, [c.value for c in ws[1]].index(HEADERS[1]) + 1).value == final
    wb.close(); assert src.read_bytes() == before
    for source, factory in ((MOS + ' другой', 'Завод А'), (MOS, 'Завод Б')):
        got = OperatorExpertAnonymizer(store.snapshot()).anonymize(source, '00001', factory)
        assert not got['knowledge'].get('exact_applied') and got['status'] == 'ЖЁЛТЫЙ'


def test_conflicts_require_individual_approval(store, tmp_path):
    decide(store); decide(store, code='00002')
    path = workbook(tmp_path, [['00001', MOS, MOS + ' новое', 'Завод А'], ['00002', MOS, MOS + ' другое', 'Завод А']])
    plan = preview(path, store)
    assert plan['conflicts'] == 2
    key = case_key('00001', MOS, 'Завод А')
    report = commit_selection(path, store, plan, [key])
    assert report['accepted'] == report['skipped_conflicts'] == 1
    assert store.snapshot().entries[key]['value'] == MOS + ' новое'
    assert store.snapshot().entries[case_key('00002', MOS, 'Завод А')]['value'] == MOS


def test_disabled_identical_result_requires_approval(store, tmp_path):
    decide(store); key = case_key('00001', MOS, 'Завод А'); store.control(key, True, 'Ошибочное')
    path = workbook(tmp_path); plan = preview(path, store)
    assert plan['conflicts'] == 1 and plan['repeated'] == 0
    assert commit_selection(path, store, plan)['accepted'] == 0
    assert store.snapshot().entries[key]['status'] == 'DISABLED'
    assert commit_selection(path, store, preview(path, store), [key])['accepted'] == 1
    assert store.snapshot().entries[key]['status'] == 'ACTIVE'


def test_inconsistent_duplicates_skipped_identical_collapsed(store, tmp_path):
    path = workbook(tmp_path, [['1', MOS, MOS, ''], ['1', MOS, MOS + ' другое', ''],
                               ['2', MOS, MOS, ''], ['2', MOS, MOS, '']])
    plan = preview(path, store)
    assert plan['new'] == 1 and plan['repeated'] == 1 and len(plan['issues']) == 2
    assert commit_selection(path, store, plan)['accepted'] == 1
    assert case_key('1', MOS) not in store.snapshot().entries


@pytest.mark.parametrize('column', range(4))
def test_formula_in_any_mapped_field_rejected(store, tmp_path, column):
    row = ['00001', MOS, MOS, 'Завод А']; row[column] = '=1+1'
    path = workbook(tmp_path, [row]); plan = preview(path, store)
    assert not plan['ready'] and len(plan['issues']) == 1


def test_missing_original_and_protected_loss_skipped(store, tmp_path):
    path = workbook(tmp_path, [['1', '', MOS, ''], ['2', MOS + ' IP66', MOS, '']])
    plan = preview(path, store)
    assert plan['ready'] == 0 and len(plan['issues']) == 2
    assert not store.snapshot().entries


def test_mapping_handles_reversed_columns_and_title_rows(store, tmp_path):
    path = tmp_path / 'reverse.xlsx'; wb = Workbook(); ws = wb.active
    ws.append(['Отчёт', '2026']); ws.append(['Обезличенное наименование', 'Наименование', 'Код Автодокс'])
    ws.append([MOS, MOS, '00001']); wb.save(path); wb.close()
    candidates = header_candidates(path); candidate = candidates[best_candidate(candidates)]
    assert candidate['row'] == 2
    mapping = guess_mapping(candidate['headers']); assert mapping['source'] == 1 and mapping['final'] == 0
    assert preview_selection(path, store, candidate, mapping)['new'] == 1


@pytest.mark.parametrize('change', ['file', 'base', 'mapping', 'store'])
def test_stale_preview_rejected_without_writes(store, tmp_path, change):
    path = workbook(tmp_path); plan = preview(path, store)
    if change == 'file':
        workbook(tmp_path, [['00001', MOS, MOS + ' изменено', 'Завод А']])
    elif change == 'base':
        decide(store, code='NEW')
    elif change == 'mapping':
        plan['plan'][0]['final'] = MOS + ' подмена'
    else:
        store = SimpleKnowledgeStore(tmp_path / 'else-local', tmp_path / 'else-share', 'engineer')
    count = len(store.events())
    with pytest.raises(ValueError):
        commit_selection(path, store, plan)
    assert len(store.events()) == count


def test_xls_formula_rejected_instead_of_using_cached_value(store, tmp_path):
    import xlwt
    path = workbook(tmp_path, [['00001', MOS, xlwt.Formula('"Клапан"'), '']], suffix='.xls')
    original = path.read_bytes()
    with pytest.raises(ValueError, match='формулы'):
        preview(path, store)
    assert path.read_bytes() == original and not store.snapshot().entries


def test_missing_or_overlapping_mapping_rejected(store, tmp_path):
    path = workbook(tmp_path); candidate = header_candidates(path)[0]
    for mapping in ({'source': None, 'final': 2}, {'source': 1, 'final': 1}, {'source': 1, 'final': 2, 'code': 1}):
        with pytest.raises(ValueError):
            preview_selection(path, store, candidate, mapping)


def test_concurrent_engineer_decision_cannot_be_silently_replaced(store, tmp_path, monkeypatch):
    other = SimpleKnowledgeStore(tmp_path / 'other', store.shared, 'engineer-two')
    path = workbook(tmp_path); plan = preview(path, store)
    original_emit = store.emit
    def emit(payload):
        decide(other, final=MOS + ' одновременная правка')
        other.sync(auto_backup=False)
        return original_emit(payload)
    monkeypatch.setattr(store, 'emit', emit)
    report = commit_selection(path, store, plan)
    entry = store.snapshot().entries[case_key('00001', MOS, 'Завод А')]
    assert entry['status'] == 'DISPUTED' and report['issues']
    assert not OperatorExpertAnonymizer(store.snapshot()).anonymize(MOS, '00001', 'Завод А')['knowledge'].get('exact_applied')


def test_bulk_import_does_not_rescan_base_per_row(store, tmp_path, monkeypatch):
    path = workbook(tmp_path, [[str(i), MOS, MOS, ''] for i in range(100)])
    plan = preview(path, store)
    counts = {'events': 0, 'snapshots': 0}
    original_events, original_snapshot = store.events, store.snapshot
    def events():
        counts['events'] += 1
        return original_events()
    def snapshot():
        counts['snapshots'] += 1
        return original_snapshot()
    monkeypatch.setattr(store, 'events', events); monkeypatch.setattr(store, 'snapshot', snapshot)
    assert commit_selection(path, store, plan)['accepted'] == 100
    assert counts['events'] < 12 and counts['snapshots'] < 5


def test_imported_decision_can_be_exported_edited_and_disabled(store, tmp_path):
    from excel.knowledge_edit_io import export_editable, commit_editable
    path = workbook(tmp_path); commit_selection(path, store, preview(path, store))
    export = tmp_path / 'База.xlsx'; export_editable(export, store)
    wb = load_workbook(export); ws = wb['Управление']; cols = {c.value: c.column for c in ws[1]}
    key = case_key('00001', MOS, 'Завод А')
    n = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, cols['Ключ']).value == key)
    ws.cell(n, cols['Новое решение'], MOS + ' исправлено')
    ws.cell(n, cols['Действие'], 'ИЗМЕНИТЬ'); ws.cell(n, cols['Причина изменения'], 'Проверено')
    wb.save(export); wb.close()
    assert commit_editable(export, store)['applied'] == 1
    assert store.snapshot().entries[key]['value'] == MOS + ' исправлено'
    export_editable(export, store); wb = load_workbook(export); ws = wb['Управление']
    ws.cell(n, cols['Действие'], 'ОТКЛЮЧИТЬ'); ws.cell(n, cols['Причина изменения'], 'Ошибка в выборке')
    wb.save(export); wb.close()
    assert commit_editable(export, store)['applied'] == 1
    assert store.snapshot().entries[key]['status'] == 'DISABLED'
